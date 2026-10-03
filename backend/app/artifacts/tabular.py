"""Bounded table decoding and formula-safe spreadsheet serialization."""

from __future__ import annotations

import csv
import io
import json
import re
import math
from datetime import date, datetime, time
from collections.abc import Iterable
from decimal import Decimal
from pathlib import PurePosixPath
from typing import Any

from openpyxl import Workbook, load_workbook

csv.field_size_limit(25 * 1024 * 1024)

_NUMBER = re.compile(r"^-?(?:0|[1-9][0-9]*)(?:\.[0-9]+)?(?:[eE][+-]?[0-9]+)?$")
_FORMULA_PREFIX = re.compile(r"^[\s\u0000-\u0020]*[=+\-@]")
_MAX_COLUMNS = 256
_MAX_CELL_CHARS = 64_000


def formula_safe_text(value: str) -> str:
    """Neutralize spreadsheet expressions without changing ordinary numbers."""
    if _FORMULA_PREFIX.match(value) and not _NUMBER.fullmatch(value.strip()):
        return "'" + value
    return value


def _string(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    return str(value)


def _normal_value(value: Any) -> Any:
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if isinstance(value, Decimal) and not value.is_finite():
        return None
    return value


def _reject_nonfinite(value: str) -> None:
    raise ValueError(f"non-finite JSON number is not supported: {value}")


def _validate_table(columns: list[str], rows: list[list[Any]]) -> None:
    if len(columns) > _MAX_COLUMNS or len(set(columns)) != len(columns):
        raise ValueError("table headers are duplicated or exceed safe bounds")
    if any(len(column) > 1_024 for column in columns):
        raise ValueError("table header exceeds safe bounds")
    if any(len(_string(cell)) > _MAX_CELL_CHARS for row in rows for cell in row):
        raise ValueError("table cell exceeds safe bounds")


def decode_table(
    content: bytes, name: str, *, max_rows: int = 250_000
) -> tuple[list[str], list[list[Any]]]:
    suffix = PurePosixPath(name).suffix.lower()
    if suffix == ".csv" or content.startswith(b"\xef\xbb\xbf"):
        text = content.decode("utf-8-sig", errors="strict")
        reader = csv.reader(io.StringIO(text, newline=""))
        try:
            columns = next(reader)
        except StopIteration:
            return [], []
        csv_rows: list[list[Any]] = []
        for row in reader:
            if len(csv_rows) >= max_rows:
                break
            csv_rows.append(
                row[: len(columns)] + [""] * max(0, len(columns) - len(row))
            )
        _validate_table(columns, csv_rows)
        return columns, csv_rows
    if suffix in {".xlsx", ".xlsm"} or content.startswith(b"PK\x03\x04"):
        from app.sources.files import _check_xlsx_archive

        _check_xlsx_archive(content)
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        try:
            sheet = workbook.active
            if sheet is None:
                return [], []
            iterator = sheet.iter_rows(values_only=True)
            header = next(iterator, ())
            columns = [_string(value) for value in header]
            rows: list[list[Any]] = []
            for row_values in iterator:
                if len(rows) >= max_rows:
                    break
                rows.append(
                    [_normal_value(value) for value in row_values[: len(columns)]]
                )
            _validate_table(columns, rows)
            return columns, rows
        finally:
            workbook.close()
    if suffix == ".parquet":
        try:
            import pyarrow.parquet as parquet  # type: ignore[import-untyped]
        except ImportError as error:
            raise ValueError("Parquet support is unavailable") from error
        parquet_file = parquet.ParquetFile(io.BytesIO(content))
        metadata = parquet_file.metadata
        total_uncompressed = sum(
            metadata.row_group(index).total_byte_size
            for index in range(metadata.num_row_groups)
        )
        if (
            metadata.num_columns > 256
            or metadata.num_row_groups > 10_000
            or metadata.num_rows > 10_000_000
            or total_uncompressed > 256 * 1024 * 1024
            or any(
                metadata.row_group(index).total_byte_size > 128 * 1024 * 1024
                for index in range(metadata.num_row_groups)
            )
        ):
            raise ValueError("Parquet metadata exceeds safe read bounds")
        columns = list(parquet_file.schema.names)
        if len(columns) > _MAX_COLUMNS or len(set(columns)) != len(columns):
            raise ValueError(
                "Parquet schema headers are duplicated or exceed safe bounds"
            )
        parquet_rows: list[list[Any]] = []
        for batch in parquet_file.iter_batches(batch_size=8192):
            for row_value in batch.to_pylist():
                parquet_rows.append(
                    [_normal_value(row_value.get(column)) for column in columns]
                )
                if len(parquet_rows) >= max_rows:
                    break
            if len(parquet_rows) >= max_rows:
                break
        _validate_table(columns, parquet_rows)
        return columns, parquet_rows
    if suffix == ".json" or content.lstrip().startswith((b"[", b"{")):
        value = json.loads(
            content,
            parse_float=Decimal,
            parse_constant=_reject_nonfinite,
        )
        if (
            isinstance(value, dict)
            and isinstance(value.get("columns"), list)
            and isinstance(value.get("rows"), list)
        ):
            columns = [_string(item) for item in value["columns"]]
            rows = value["rows"][:max_rows]
            result = [_row_values(row, columns) for row in rows]
            _validate_table(columns, result)
            return columns, result
        if (
            isinstance(value, list)
            and value
            and all(isinstance(row, dict) for row in value)
        ):
            columns = list(dict.fromkeys(key for row in value for key in row))
            result = [
                [row.get(column) for column in columns] for row in value[:max_rows]
            ]
            _validate_table(columns, result)
            return columns, result
    raise ValueError("artifact is not a supported tabular result")


def _row_values(row: Any, columns: list[str]) -> list[Any]:
    if isinstance(row, dict):
        return [row.get(column) for column in columns]
    if isinstance(row, list):
        return (row + [None] * len(columns))[: len(columns)]
    raise ValueError("tabular JSON rows must be objects or arrays")


def safe_csv(columns: list[str], rows: Iterable[Iterable[Any]]) -> bytes:
    stream = io.StringIO(newline="")
    writer = csv.writer(stream, lineterminator="\r\n")
    writer.writerow([formula_safe_text(_string(value)) for value in columns])
    for row in rows:
        writer.writerow(
            [formula_safe_text(_string(_normal_value(value))) for value in row]
        )
    return stream.getvalue().encode("utf-8")


def safe_csv_export(
    original: bytes, columns: list[str], rows: Iterable[Iterable[Any]]
) -> bytes:
    """Keep canonical CSV bytes unless formula escaping changes a cell."""
    materialized_rows = [list(row) for row in rows]
    unchanged = all(
        formula_safe_text(_string(value)) == _string(value) for value in columns
    ) and all(
        formula_safe_text(_string(_normal_value(value)))
        == _string(_normal_value(value))
        for row in materialized_rows
        for value in row
    )
    if unchanged:
        return original
    return safe_csv(columns, materialized_rows)


def _xlsx_value(value: Any) -> str | int | float | bool | None:
    value = _normal_value(value)
    if value is None or isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        # Excel stores only 15 significant decimal digits. Preserve large ints as text.
        return (
            value
            if not isinstance(value, int) or len(str(abs(value))) <= 15
            else str(value)
        )
    if isinstance(value, Decimal):
        return str(value)
    return formula_safe_text(_string(value))


def safe_xlsx(columns: list[str], rows: Iterable[Iterable[Any]]) -> bytes:
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Results")
    sheet.append([formula_safe_text(_string(value)) for value in columns])
    for row in rows:
        sheet.append([_xlsx_value(value) for value in row])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def safe_parquet(columns: list[str], rows: list[list[Any]]) -> bytes:
    try:
        import pyarrow as arrow
        import pyarrow.parquet as parquet
    except ImportError as error:
        raise ValueError("Parquet support is unavailable") from error
    arrays: dict[str, list[Any]] = {}
    for index, column in enumerate(columns):
        values = [row[index] if index < len(row) else None for row in rows]
        arrays[column] = values
    table = arrow.table(arrays)
    output = io.BytesIO()
    parquet.write_table(table, output)
    return output.getvalue()
