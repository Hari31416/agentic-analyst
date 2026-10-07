"""Bounded local extractors for documents supported by the ingestion pipeline."""

from __future__ import annotations

import csv
import io
import math
import re
import shutil
import subprocess
import time
import zipfile
from collections.abc import Iterable
from html.parser import HTMLParser
from pathlib import PurePosixPath
from typing import Any

from pypdf import PdfReader

from app.ingestion.pdf_text import (
    PdfTextLimitExceeded,
    bounded_plain_text,
    legacy_font_names,
    normalize_pdf_whitespace,
)

from app.sources.documents import (
    DocumentIngestionError,
    ExtractedBlock,
    _block,
    _validate_docx_archive,
    _xml_text_guard,
    MAX_EXTRACTED_BYTES,
    MAX_PDF_PAGES,
    MAX_ARCHIVE_BYTES,
    MAX_ARCHIVE_MEMBERS,
    MAX_MEMBER_BYTES,
    MAX_RATIO,
)

MAX_OCR_PAGES = 100
MAX_HTML_NODES = 200_000
MAX_PPTX_SLIDES = 2_000
MAX_RENDER_EDGE = 6_000
MAX_RENDER_PIXELS = 24_000_000
_DOCTYPE_RE = re.compile(rb"<!\s*(?:DOCTYPE|ENTITY)\b", re.IGNORECASE)


def _safe_archive(content: bytes, extension: str) -> set[str]:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > MAX_ARCHIVE_MEMBERS:
                raise DocumentIngestionError(
                    "archive_too_large",
                    f"{extension.upper()} archive contains too many members.",
                )
            names: set[str] = set()
            expanded = 0
            for member in members:
                name = member.filename.replace("\\", "/")
                path = PurePosixPath(name)
                if path.is_absolute() or ".." in path.parts or name in names:
                    raise DocumentIngestionError(
                        "unsafe_archive",
                        f"{extension.upper()} archive contains an unsafe path.",
                    )
                names.add(name)
                if member.file_size > MAX_MEMBER_BYTES:
                    raise DocumentIngestionError(
                        "archive_member_too_large",
                        "Archive member exceeds the size limit.",
                    )
                expanded += member.file_size
                if member.file_size > max(1, member.compress_size) * MAX_RATIO:
                    raise DocumentIngestionError(
                        "archive_ratio_exceeded",
                        "Archive compression ratio exceeds the limit.",
                    )
                lowered = name.lower()
                if any(
                    part in lowered
                    for part in ("vbaproject.bin", "activex/", "embeddings/")
                ):
                    raise DocumentIngestionError(
                        "active_content_rejected",
                        "Macros and embedded active content are not supported.",
                    )
                if lowered.endswith((".xml", ".rels")):
                    raw = archive.read(member)
                    if _DOCTYPE_RE.search(raw):
                        raise DocumentIngestionError(
                            "unsafe_xml",
                            "Office document contains a prohibited XML declaration.",
                        )
                    if b"macroenabled" in raw.lower():
                        raise DocumentIngestionError(
                            "active_content_rejected",
                            "Macro-enabled Office documents are not supported.",
                        )
            if expanded > MAX_ARCHIVE_BYTES:
                raise DocumentIngestionError(
                    "archive_too_large",
                    "Expanded archive content exceeds the size limit.",
                )
            return names
    except DocumentIngestionError:
        raise
    except (OSError, zipfile.BadZipFile, RuntimeError) as error:
        raise DocumentIngestionError(
            "invalid_archive", f"The {extension.upper()} package is invalid."
        ) from error


def _check_size(blocks: Iterable[ExtractedBlock]) -> None:
    total = sum(len(block.text.encode("utf-8")) for block in blocks)
    if total > MAX_EXTRACTED_BYTES:
        raise DocumentIngestionError(
            "extracted_text_too_large",
            "Extracted document text exceeds the processing limit.",
        )


def _pdfium_page_png(content: bytes, page_number: int, scale: float = 2.0) -> bytes:
    try:
        import pypdfium2 as pdfium  # type: ignore[import-untyped]
    except ImportError as error:
        raise DocumentIngestionError(
            "ocr_renderer_unavailable",
            "Scanned PDF OCR requires the optional pypdfium2 renderer.",
        ) from error
    document = pdfium.PdfDocument(content)
    page = None
    bitmap = None
    image = None
    try:
        page = document[page_number - 1]
        width_points, height_points = page.get_size()
        if (
            not math.isfinite(width_points)
            or not math.isfinite(height_points)
            or width_points <= 0
            or height_points <= 0
        ):
            raise DocumentIngestionError(
                "invalid_pdf_page_size", "PDF page has invalid dimensions."
            )
        bounded_scale = min(
            scale,
            MAX_RENDER_EDGE / max(width_points, height_points),
            math.sqrt(MAX_RENDER_PIXELS / (width_points * height_points)),
        )
        if bounded_scale < 0.1:
            raise DocumentIngestionError(
                "pdf_render_dimensions_exceeded",
                "PDF page dimensions exceed the safe local rendering limit.",
            )
        bitmap = page.render(scale=bounded_scale)
        image = bitmap.to_pil()
        output = io.BytesIO()
        image.save(output, format="PNG")
        return output.getvalue()
    finally:
        if image is not None:
            image.close()
        if bitmap is not None:
            bitmap.close()
        if page is not None:
            page.close()
        document.close()


def _tesseract_languages(executable: str) -> set[str]:
    try:
        result = subprocess.run(
            [executable, "--list-langs"],
            capture_output=True,
            text=True,
            timeout=5,
            check=False,
        )
    except (OSError, subprocess.TimeoutExpired):
        return set()
    if result.returncode != 0:
        return set()
    return {line.strip() for line in result.stdout.splitlines()[1:] if line.strip()}


def _check_tesseract_availability(languages: str) -> tuple[str | None, str | None]:
    executable = shutil.which("tesseract")
    if executable is None:
        return None, "ocr_unavailable:tesseract_missing"
    requested = {item.strip() for item in languages.split("+") if item.strip()}
    available = _tesseract_languages(executable)
    missing = sorted(requested - available)
    if missing:
        return None, "ocr_unavailable:missing_language_assets:" + "+".join(missing)
    return executable, None


def _run_tesseract(
    png: bytes, languages: str, timeout: int, executable: str | None = None
) -> tuple[str, str | None, dict[str, Any]]:
    if executable is None:
        executable, issue = _check_tesseract_availability(languages)
        if issue:
            return "", issue, {}
    assert executable is not None
    requested = {item.strip() for item in languages.split("+") if item.strip()}
    try:
        deadline = time.monotonic() + timeout
        rotation = 0
        try:
            osd = subprocess.run(
                [executable, "stdin", "stdout", "-l", "osd", "--psm", "0"],
                input=png,
                capture_output=True,
                timeout=min(4, max(0.1, deadline - time.monotonic())),
                check=False,
            )
            if osd.returncode == 0:
                match = re.search(rb"Rotate:\s*(0|90|180|270)\b", osd.stdout)
                if match:
                    rotation = int(match.group(1))
        except (OSError, subprocess.TimeoutExpired):
            pass
        if rotation:
            from PIL import Image

            with Image.open(io.BytesIO(png)) as image:
                corrected = image.rotate(-rotation, expand=True)
                output = io.BytesIO()
                corrected.save(output, format="PNG")
                corrected.close()
                png = output.getvalue()
        remaining = deadline - time.monotonic()
        if remaining <= 0:
            return "", "ocr_timeout", {}
        result = subprocess.run(
            [
                executable,
                "stdin",
                "stdout",
                "-l",
                "+".join(sorted(requested)),
                "--psm",
                "3",
                "tsv",
            ],
            input=png,
            capture_output=True,
            timeout=remaining,
            check=False,
        )
    except DocumentIngestionError as error:
        return "", error.code, {}
    except subprocess.TimeoutExpired:
        return "", "ocr_timeout", {}
    except OSError:
        return "", "ocr_unavailable:tesseract_execution_failed", {}
    if result.returncode != 0:
        return "", "ocr_failed", {}
    words: list[str] = []
    boxes: list[dict[str, int]] = []
    confidences: list[float] = []
    try:
        reader = csv.DictReader(
            io.StringIO(result.stdout.decode("utf-8", errors="replace")),
            delimiter="\t",
        )
        for row in reader:
            word = (row.get("text") or "").strip()
            if not word:
                continue
            words.append(word)
            try:
                confidence = float(row.get("conf", "-1"))
                if confidence >= 0:
                    confidences.append(confidence)
                boxes.append(
                    {
                        "left": int(row["left"]),
                        "top": int(row["top"]),
                        "width": int(row["width"]),
                        "height": int(row["height"]),
                    }
                )
            except (KeyError, TypeError, ValueError):
                continue
    except csv.Error:
        return "", "ocr_failed:invalid_tsv", {}
    metadata: dict[str, Any] = {
        "ocr_method": "tesseract_tsv",
        "ocr_rotation_degrees": rotation,
        "ocr_word_count": len(words),
        "ocr_word_boxes": boxes[:2_000],
    }
    if confidences:
        metadata["ocr_confidence"] = round(sum(confidences) / len(confidences), 2)
        metadata["ocr_confidence_type"] = "mean_tesseract_word_confidence_percent"
    return " ".join(words), None, metadata


def _ocr_page(
    content: bytes, page_number: int, languages: str, timeout: int
) -> tuple[str, str | None, dict[str, Any]]:
    executable, issue = _check_tesseract_availability(languages)
    if issue:
        return "", issue, {}
    try:
        png = _pdfium_page_png(content, page_number)
    except DocumentIngestionError as error:
        return "", error.code, {}
    return _run_tesseract(
        png, languages=languages, timeout=timeout, executable=executable
    )


def _ocr_image(
    content: bytes, languages: str, timeout: int
) -> tuple[str, str | None, dict[str, Any]]:
    executable, issue = _check_tesseract_availability(languages)
    if issue:
        return "", issue, {}
    try:
        from PIL import Image

        with Image.open(io.BytesIO(content)) as img:
            if img.format == "PNG":
                png = content
            else:
                out = io.BytesIO()
                img.convert("RGB").save(out, format="PNG")
                png = out.getvalue()
    except Exception:
        return "", "invalid_image", {}
    return _run_tesseract(
        png, languages=languages, timeout=timeout, executable=executable
    )


def extract_image(
    filename: str,
    content: bytes,
    *,
    ocr_enabled: bool = True,
    ocr_languages: str = "eng+hin",
    ocr_timeout_seconds: int = 20,
) -> tuple[list[ExtractedBlock], list[str], list[str]]:
    try:
        from PIL import Image

        with Image.open(io.BytesIO(content)) as img:
            width, height = img.size
            img_format = (img.format or "PNG").lower()
            if width <= 0 or height <= 0 or width * height > MAX_RENDER_PIXELS:
                raise DocumentIngestionError(
                    "invalid_image_size",
                    "Image dimensions are invalid or exceed limits.",
                )
    except DocumentIngestionError:
        raise
    except Exception as error:
        raise DocumentIngestionError(
            "invalid_image", "The image could not be read safely."
        ) from error

    warnings: list[str] = []
    text = ""
    location: dict[str, Any] = {
        "type": "image",
        "filename": filename,
        "image_format": img_format,
        "image_width": width,
        "image_height": height,
        "image_size_bytes": len(content),
    }
    if ocr_enabled:
        recognized, issue, ocr_metadata = _ocr_image(
            content, ocr_languages, ocr_timeout_seconds
        )
        if issue:
            warnings.append(f"image_ocr:{issue}")
        if recognized.strip():
            text = recognized.strip()
            location["extractor"] = "ocr"
            location.update(ocr_metadata)
    if not text:
        text = f"[Image: {filename}]"
        location["extractor"] = "image"

    block = _block("image", text, None, location, image_bytes=content)
    return [block], warnings, [block.language]


def extract_pdf(
    content: bytes,
    *,
    ocr_enabled: bool = True,
    ocr_languages: str = "eng+hin",
    ocr_timeout_seconds: int = 20,
) -> tuple[list[ExtractedBlock], list[str], list[str]]:
    try:
        reader = PdfReader(io.BytesIO(content), strict=True)
        if reader.is_encrypted:
            raise DocumentIngestionError(
                "encrypted_pdf", "Encrypted PDFs are not supported."
            )
        if len(reader.pages) > MAX_PDF_PAGES:
            raise DocumentIngestionError(
                "pdf_page_limit", "PDF has too many pages to process."
            )
        legacy_fonts_by_page = {
            page_number: fonts
            for page_number, page in enumerate(reader.pages, start=1)
            if (fonts := legacy_font_names(page))
        }
        legacy_page_numbers = sorted(legacy_fonts_by_page)
        if legacy_page_numbers and not ocr_enabled:
            raise DocumentIngestionError(
                "legacy_font_ocr_required",
                "PDF pages use non-Unicode legacy fonts and require OCR. Enable OCR or upload an OCR-ready copy.",
            )
        if len(legacy_page_numbers) > MAX_OCR_PAGES:
            page_sample = ",".join(str(page) for page in legacy_page_numbers[:20])
            raise DocumentIngestionError(
                "ocr_page_limit",
                f"PDF has {len(legacy_page_numbers)} pages using non-Unicode legacy fonts; local OCR is limited to {MAX_OCR_PAGES} pages per document. Split the PDF into smaller parts or upload the relevant chapters. First affected pages: {page_sample}.",
            )
        blocks: list[ExtractedBlock] = []
        warnings: list[str] = []
        empty_pages: list[int] = []
        legacy_pages: list[int] = []
        legacy_ocr_failures: list[tuple[int, str]] = []
        total = 0
        ocr_attempts = 0
        toc_entries: dict[int, list[tuple[int, str]]] = {}

        def future_legacy_pages(current_page: int) -> int:
            return sum(page > current_page for page in legacy_page_numbers)

        def visit_outlines(items: list[Any], level: int = 1) -> None:
            for item in items:
                if isinstance(item, list):
                    visit_outlines(item, level + 1)
                    continue
                try:
                    page_index = reader.get_destination_page_number(item)
                    if page_index is None:
                        continue
                    title = str(item.title).strip()
                    if title:
                        toc_entries.setdefault(page_index + 1, []).append(
                            (level, title)
                        )
                except Exception:
                    continue

        try:
            visit_outlines(reader.outline)
        except Exception:
            pass
        embedded_image_count = 0
        MAX_EMBEDDED_IMAGES = 100

        def extract_page_images(p_num: int, p_obj: Any, p_heading: str | None) -> None:
            nonlocal embedded_image_count
            if embedded_image_count >= MAX_EMBEDDED_IMAGES:
                return
            for img in getattr(p_obj, "images", []):
                if embedded_image_count >= MAX_EMBEDDED_IMAGES:
                    break
                try:
                    img_data = img.data
                    if len(img_data) < 100 or len(img_data) > MAX_MEMBER_BYTES:
                        continue
                    from PIL import Image

                    with Image.open(io.BytesIO(img_data)) as pil_img:
                        w, h = pil_img.size
                        fmt = (pil_img.format or "PNG").lower()
                        if w < 32 or h < 32:
                            continue
                except Exception:
                    continue
                embedded_image_count += 1
                img_name = getattr(img, "name", f"image_{embedded_image_count}")
                img_text = f"[Image on page {p_num}: {img_name}]"
                img_location = {
                    "type": "image",
                    "page": p_num,
                    "image_name": img_name,
                    "image_format": fmt,
                    "image_width": w,
                    "image_height": h,
                    "image_size_bytes": len(img_data),
                }
                blocks.append(
                    _block(
                        "image",
                        img_text,
                        p_heading,
                        img_location,
                        image_bytes=img_data,
                    )
                )

        for page_number, page in enumerate(reader.pages, start=1):
            legacy_fonts = legacy_fonts_by_page.get(page_number, [])
            is_legacy = bool(legacy_fonts)
            if is_legacy:
                legacy_pages.append(page_number)
                # Legacy font character codes are not Unicode. Never publish
                # pypdf's visually plausible but semantically unrelated text.
                text = ""
            else:
                try:
                    remaining = MAX_EXTRACTED_BYTES - total
                    text = normalize_pdf_whitespace(bounded_plain_text(page, remaining))
                except PdfTextLimitExceeded as error:
                    raise DocumentIngestionError(
                        "extracted_text_too_large",
                        "Extracted document text exceeds the processing limit.",
                    ) from error
                except Exception as error:
                    raise DocumentIngestionError(
                        "pdf_extraction_failed", "PDF text extraction failed."
                    ) from error
            weak = not is_legacy and (
                len(text.strip()) < 24 or text.count("\ufffd") > max(2, len(text) // 20)
            )
            source = "digital"
            confidence = None
            if (is_legacy or not text.strip() or weak) and ocr_enabled:
                ordinary_page_limit = (
                    MAX_OCR_PAGES
                    if is_legacy
                    else MAX_OCR_PAGES - future_legacy_pages(page_number)
                )
                if ocr_attempts >= ordinary_page_limit:
                    warnings.append(f"ocr_page_limit_reached:{page_number}")
                    if is_legacy:
                        legacy_ocr_failures.append(
                            (page_number, "ocr_page_limit_reached")
                        )
                else:
                    ocr_attempts += 1
                    recognized, issue, ocr_metadata = _ocr_page(
                        content, page_number, ocr_languages, ocr_timeout_seconds
                    )
                    if issue:
                        warnings.append(f"page_{page_number}:{issue}")
                        if is_legacy:
                            legacy_ocr_failures.append((page_number, issue))
                    if recognized.strip() and (
                        not text.strip() or len(recognized) > len(text)
                    ):
                        text = recognized
                        source = "ocr"
                        # This is a recognizer signal only; it is not verified accuracy.
                        confidence = ocr_metadata
                    elif is_legacy and not issue:
                        legacy_ocr_failures.append((page_number, "empty_ocr_result"))
            elif is_legacy:
                legacy_ocr_failures.append((page_number, "ocr_disabled"))
            headings = toc_entries.get(page_number, [])
            parent = " > ".join(title for _, title in headings) if headings else None
            extract_page_images(page_number, page, parent)
            if not text.strip():
                empty_pages.append(page_number)
                continue
            total += len(text.encode("utf-8"))
            if total > MAX_EXTRACTED_BYTES:
                raise DocumentIngestionError(
                    "extracted_text_too_large",
                    "Extracted document text exceeds the processing limit.",
                )
            location: dict[str, Any] = {
                "page": page_number,
                "char_start": 0,
                "char_end": len(text),
                "extractor": source,
            }
            if source == "digital":
                location["text_mode"] = "plain"
            if legacy_fonts:
                location["legacy_fonts"] = legacy_fonts
            if confidence:
                location.update(confidence)
                if "ocr_confidence" not in location:
                    location["ocr_confidence"] = "not_supplied"
                warnings.append(f"page_{page_number}:ocr_table_cells_not_verified")
                confidence_value = location.get("ocr_confidence")
                if isinstance(confidence_value, (int, float)) and confidence_value < 70:
                    warnings.append(
                        f"page_{page_number}:low_ocr_recognizer_confidence_review_required"
                    )
            if headings:
                warnings.append(f"page_{page_number}:pdf_outline_headings_applied")
            lines = text.splitlines()
            numbered_heading = re.compile(
                r"^(?:chapter|section|appendix)\s+[\w.-]+\b|^\d+(?:\.\d+){0,4}[.)]?\s+\S",
                re.IGNORECASE,
            )
            nonempty_lines = [line for line in lines if line.strip()]
            if len(nonempty_lines) > 1:
                current_heading = parent
                cursor = 0
                for line in lines:
                    stripped = line.strip()
                    if not stripped:
                        cursor += len(line) + 1
                        continue
                    line_location = {
                        **location,
                        "char_start": cursor,
                        "char_end": cursor + len(line),
                    }
                    if numbered_heading.search(stripped) and len(stripped) <= 160:
                        current_heading = (
                            f"{parent} > {stripped}" if parent else stripped
                        )
                        blocks.append(
                            _block("heading", stripped, current_heading, line_location)
                        )
                    else:
                        blocks.append(
                            _block(
                                "paragraph", stripped, current_heading, line_location
                            )
                        )
                    cursor += len(line) + 1
            else:
                blocks.append(_block("page_text", text, parent, location))
        if legacy_ocr_failures:
            failures = sorted(set(legacy_ocr_failures))
            page_sample = ", ".join(
                f"{page} ({reason})" for page, reason in failures[:20]
            )
            suffix = "" if len(failures) <= 20 else ", ..."
            raise DocumentIngestionError(
                "legacy_font_ocr_failed",
                f"OCR failed for {len(failures)} legacy-font page(s). Check local OCR availability and retry with working Hindi language assets. Failed pages: {page_sample}{suffix}.",
            )
        if empty_pages:
            warnings.append(
                "ocr_needed_pages:" + ",".join(str(page) for page in empty_pages[:100])
            )
            warnings.append(
                "Some PDF pages contain no extractable text and may require OCR."
            )
        if legacy_pages:
            warnings.append(
                "legacy_font_pages_detected:"
                + ",".join(str(page) for page in legacy_pages[:100])
            )
        if not blocks:
            warnings.append("ocr_needed")
        return (
            blocks,
            list(dict.fromkeys(warnings)),
            sorted({block.language for block in blocks}),
        )
    except DocumentIngestionError:
        raise
    except Exception as error:
        raise DocumentIngestionError(
            "invalid_pdf", "The PDF could not be read safely."
        ) from error


def extract_text(
    filename: str, content: bytes
) -> tuple[list[ExtractedBlock], list[str], list[str]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise DocumentIngestionError(
            "invalid_text_encoding",
            "Text, Markdown, and HTML uploads must use UTF-8 encoding.",
        ) from error
    if "\x00" in text:
        raise DocumentIngestionError(
            "invalid_text_content", "Text document contains prohibited NUL characters."
        )
    kind = PurePosixPath(filename).suffix.lower()
    blocks: list[ExtractedBlock] = []
    if kind in {".txt", ".md"}:
        heading_stack: list[str] = []
        offset = 0
        for line_index, line in enumerate(text.splitlines(keepends=True)):
            match = (
                re.match(r"^(#{1,6})\s+(.+?)\s*#*\s*$", line.rstrip("\r\n"))
                if kind == ".md"
                else None
            )
            content_line = line.rstrip("\r\n")
            if not content_line.strip():
                offset += len(line)
                continue
            location = {
                "line": line_index + 1,
                "char_start": offset,
                "char_end": offset + len(content_line),
            }
            if match:
                level = len(match.group(1))
                heading_stack = heading_stack[: level - 1]
                heading_stack.append(match.group(2))
                blocks.append(
                    _block(
                        "heading", match.group(2), " > ".join(heading_stack), location
                    )
                )
            else:
                blocks.append(
                    _block(
                        "paragraph",
                        content_line,
                        " > ".join(heading_stack) or None,
                        location,
                    )
                )
            offset += len(line)
    else:
        parser = _SafeHtmlParser()
        try:
            parser.feed(text)
            parser.close()
            parser._flush()
        except Exception as error:
            raise DocumentIngestionError(
                "invalid_html", "HTML document could not be parsed safely."
            ) from error
        if parser.nodes > MAX_HTML_NODES:
            raise DocumentIngestionError(
                "html_node_limit", "HTML document contains too many elements."
            )
        if parser.suppressed:
            warnings = ["HTML scripts, styles, and embedded content were omitted."]
        else:
            warnings = []
        if parser.images_omitted:
            warnings.append("HTML embedded images were not OCR processed.")
        blocks.extend(parser.blocks)
    _check_size(blocks)
    if not blocks:
        raise DocumentIngestionError(
            "empty_document", "Document contains no extractable text."
        )
    return (
        blocks,
        locals().get("warnings", []),
        sorted({block.language for block in blocks}),
    )


class _SafeHtmlParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.blocks: list[ExtractedBlock] = []
        self.stack: list[str] = []
        self.active: list[str] = []
        self.parts: list[str] = []
        self.suppressed = False
        self.images_omitted = False
        self.nodes = 0
        self.table = 0
        self.row = 0
        self.cells: list[str] = []
        self.cell = ""
        self.in_cell = False

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        self.nodes += 1
        if tag in {"script", "style", "iframe", "object", "embed", "svg", "noscript"}:
            self.suppressed = True
            self.active.append("!" + tag)
            return
        if self.active and self.active[-1].startswith("!"):
            return
        heading = re.fullmatch(r"h([1-6])", tag)
        if heading:
            self._flush()
            self.stack = self.stack[: int(heading.group(1)) - 1]
            self.active.append("heading")
        elif tag in {
            "p",
            "li",
            "blockquote",
            "pre",
            "div",
            "body",
            "section",
            "article",
            "main",
            "address",
            "figure",
            "figcaption",
        }:
            self._flush()
            self.active.append("paragraph")
        elif tag == "img":
            self.images_omitted = True
        elif tag == "table":
            self.table += 1
            self.row = 0
        elif tag == "tr":
            self.cells = []
        elif tag in {"td", "th"}:
            self.cell = ""
            self.in_cell = True

    def handle_endtag(self, tag: str) -> None:
        if self.active and self.active[-1] == "!" + tag:
            self.active.pop()
            return
        if (
            self.active
            and self.active[-1] in {"paragraph", "heading"}
            and tag
            in {
                "p",
                "li",
                "blockquote",
                "pre",
                "div",
                "body",
                "section",
                "article",
                "main",
                "address",
                "figure",
                "figcaption",
                "h1",
                "h2",
                "h3",
                "h4",
                "h5",
                "h6",
            }
        ):
            self._flush()
            self.active.pop()
        elif tag in {"td", "th"}:
            self.cells.append(re.sub(r"\s+", " ", self.cell).strip())
            self.cell = ""
            self.in_cell = False
        elif tag == "tr" and self.cells:
            values = list(self.cells)
            row_number = self.row
            rendered = " | ".join(
                values
                if row_number == 0
                else [f"{i + 1}: {value}" for i, value in enumerate(values)]
            )
            self.blocks.append(
                _block(
                    "table_row",
                    rendered,
                    " > ".join(self.stack) or None,
                    {
                        "table": self.table - 1,
                        "row": row_number,
                        "cell_count": len(values),
                        "cells": values,
                        "html_block": len(self.blocks),
                        "cell_locations": [
                            {"row": row_number, "column": column}
                            for column in range(len(values))
                        ],
                    },
                )
            )
            self.row += 1

    def handle_data(self, data: str) -> None:
        if self.active and self.active[-1].startswith("!"):
            return
        if self.in_cell:
            self.cell += data
        else:
            self.parts.append(data)

    def _flush(self) -> None:
        value = re.sub(r"\s+", " ", " ".join(self.parts)).strip()
        self.parts.clear()
        if not value:
            return
        is_heading = self.active and self.active[-1] == "heading"
        if is_heading:
            self.stack.append(value)
            kind = "heading"
        else:
            kind = "paragraph"
        self.blocks.append(
            _block(
                kind,
                value,
                " > ".join(self.stack) or None,
                {"html_block": len(self.blocks)},
            )
        )


def extract_pptx(content: bytes) -> tuple[list[ExtractedBlock], list[str], list[str]]:
    names = _safe_archive(content, "pptx")
    if "[Content_Types].xml" not in names or "ppt/presentation.xml" not in names:
        raise DocumentIngestionError(
            "invalid_pptx", "The ZIP file is not a valid PPTX document."
        )
    try:
        for name in names:
            if name.endswith(".xml") and (
                name.startswith("ppt/slides/") or name == "ppt/presentation.xml"
            ):
                with zipfile.ZipFile(io.BytesIO(content)) as archive:
                    _xml_text_guard(archive.read(name))
        from pptx import Presentation

        presentation = Presentation(io.BytesIO(content))
    except ImportError as error:
        raise DocumentIngestionError(
            "pptx_extractor_unavailable",
            "PPTX ingestion requires the optional python-pptx package.",
        ) from error
    except DocumentIngestionError:
        raise
    except Exception as error:
        raise DocumentIngestionError(
            "invalid_pptx", "PPTX could not be read safely."
        ) from error
    if len(presentation.slides) > MAX_PPTX_SLIDES:
        raise DocumentIngestionError(
            "pptx_slide_limit", "PPTX contains too many slides."
        )
    blocks: list[ExtractedBlock] = []
    warnings: list[str] = []
    total = 0
    next_table_id = 0
    for slide_number, slide in enumerate(presentation.slides, start=1):
        title = slide.shapes.title.text.strip() if slide.shapes.title else ""
        heading_stack = [title] if title else []
        for shape_index, shape in enumerate(slide.shapes):
            if getattr(shape, "has_table", False):
                table = shape.table
                table_id = next_table_id
                next_table_id += 1
                headers = (
                    [
                        cell.text.replace("\n", " ").strip()
                        for cell in table.rows[0].cells
                    ]
                    if table.rows
                    else []
                )
                for row_index, row in enumerate(table.rows):
                    values = [
                        cell.text.replace("\n", " ").strip() for cell in row.cells
                    ]
                    text = " | ".join(
                        values
                        if row_index == 0
                        else [
                            f"{headers[i] if i < len(headers) and headers[i] else f'Column {i + 1}'}: {value}"
                            for i, value in enumerate(values)
                        ]
                    )
                    if text.strip(" |:"):
                        total += len(text.encode("utf-8"))
                        blocks.append(
                            _block(
                                "table_row",
                                text,
                                " > ".join(heading_stack) or None,
                                {
                                    "slide": slide_number,
                                    "shape": shape_index,
                                    "table": table_id,
                                    "row": row_index,
                                    "cell_count": len(values),
                                    "cells": values,
                                    "cell_locations": [
                                        {"row": row_index, "column": column_index}
                                        for column_index in range(len(values))
                                    ],
                                },
                            )
                        )
                continue
            if not getattr(shape, "has_text_frame", False):
                continue
            for paragraph_index, paragraph in enumerate(shape.text_frame.paragraphs):
                text = paragraph.text.strip()
                if not text:
                    continue
                total += len(text.encode("utf-8"))
                level = min(6, paragraph.level + 1)
                if shape == slide.shapes.title:
                    kind = "heading"
                elif paragraph.level > 0 or (
                    paragraph.runs and paragraph.runs[0].font.bold
                ):
                    kind = "heading"
                else:
                    kind = "paragraph"
                if kind == "heading":
                    heading_stack = heading_stack[: level - 1]
                    heading_stack.append(text)
                blocks.append(
                    _block(
                        kind,
                        text,
                        " > ".join(heading_stack) or None,
                        {
                            "slide": slide_number,
                            "shape": shape_index,
                            "paragraph": paragraph_index,
                            "reading_order": shape_index,
                        },
                    )
                )
            if getattr(shape, "shape_type", None) == 13 and getattr(
                shape, "image", None
            ):
                try:
                    img_data = shape.image.blob
                    if len(img_data) >= 100 and len(img_data) <= MAX_MEMBER_BYTES:
                        from PIL import Image

                        with Image.open(io.BytesIO(img_data)) as pil_img:
                            w, h = pil_img.size
                            fmt = (pil_img.format or "PNG").lower()
                            if w >= 32 and h >= 32:
                                img_name = f"slide_{slide_number}_image_{shape_index}"
                                blocks.append(
                                    _block(
                                        "image",
                                        f"[Image on slide {slide_number}: {img_name}]",
                                        " > ".join(heading_stack) or None,
                                        {
                                            "type": "image",
                                            "slide": slide_number,
                                            "shape": shape_index,
                                            "image_name": img_name,
                                            "image_format": fmt,
                                            "image_width": w,
                                            "image_height": h,
                                            "image_size_bytes": len(img_data),
                                        },
                                        image_bytes=img_data,
                                    )
                                )
                except Exception:
                    pass
    if total > MAX_EXTRACTED_BYTES:
        raise DocumentIngestionError(
            "extracted_text_too_large",
            "Extracted document text exceeds the processing limit.",
        )
    if not blocks:
        raise DocumentIngestionError(
            "empty_document", "PPTX contains no extractable text."
        )
    return blocks, warnings, sorted({block.language for block in blocks})


def extract_pdf_layout(
    content: bytes,
    *,
    ocr_enabled: bool = True,
    ocr_languages: str = "eng+hin",
    ocr_timeout_seconds: int = 20,
) -> tuple[list[ExtractedBlock], list[str], list[str]]:
    """Use a local geometry based extractor for PDF table candidates."""
    try:
        import pdfplumber
    except ImportError as error:
        raise DocumentIngestionError(
            "profile_unavailable",
            "The optional layout profile requires pdfplumber to be installed.",
        ) from error
    blocks, warnings, languages = extract_pdf(
        content,
        ocr_enabled=ocr_enabled,
        ocr_languages=ocr_languages,
        ocr_timeout_seconds=ocr_timeout_seconds,
    )
    layout_rows: list[ExtractedBlock] = []
    page_headings = {
        int(block.location["page"]): block.heading
        for block in blocks
        if block.kind == "heading" and "page" in block.location
    }
    try:
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            table_id = 0
            for page_number, page in enumerate(pdf.pages, start=1):
                for table in page.find_tables()[:256]:
                    rows = table.extract() or []
                    if not rows:
                        continue
                    headers = [str(cell or "").strip() for cell in rows[0]]
                    if (
                        len(rows) > 10_000
                        or max((len(row) for row in rows), default=0) > 2_000
                    ):
                        raise DocumentIngestionError(
                            "pdf_table_dimensions_exceeded",
                            "PDF table exceeds the safe row or column processing limit.",
                        )
                    for row_index, row in enumerate(rows):
                        values = [str(cell or "").strip() for cell in row]
                        rendered = (
                            values
                            if row_index == 0
                            else [
                                f"{headers[index] if index < len(headers) and headers[index] else f'Column {index + 1}'}: {value}"
                                for index, value in enumerate(values)
                            ]
                        )
                        text = " | ".join(rendered)
                        if not text.strip(" |:"):
                            continue
                        row_boxes = (
                            table.rows[row_index].cells
                            if row_index < len(table.rows)
                            else []
                        )
                        layout_rows.append(
                            _block(
                                "table_row",
                                text,
                                page_headings.get(page_number),
                                {
                                    "page": page_number,
                                    "table": table_id,
                                    "row": row_index,
                                    "cell_count": len(values),
                                    "cells": values,
                                    "cell_locations": [
                                        {
                                            "row": row_index,
                                            "column": column,
                                            "bbox": list(box) if box else None,
                                        }
                                        for column, box in enumerate(row_boxes)
                                    ],
                                    "extractor": "pdfplumber_layout",
                                    "review_required": True,
                                },
                            )
                        )
                    table_id += 1
    except Exception:
        warnings.append("layout_table_detection_failed")
    base_by_page: dict[int, list[ExtractedBlock]] = {}
    tables_by_page: dict[int, list[ExtractedBlock]] = {}
    other_blocks: list[ExtractedBlock] = []
    for block in blocks:
        page_ref = block.location.get("page")
        if isinstance(page_ref, int):
            base_by_page.setdefault(page_ref, []).append(block)
        else:
            other_blocks.append(block)
    for block in layout_rows:
        page_number = int(block.location["page"])
        tables_by_page.setdefault(page_number, []).append(block)
    ordered_blocks = list(other_blocks)
    for page_number in sorted(set(base_by_page) | set(tables_by_page)):
        ordered_blocks.extend(base_by_page.get(page_number, []))
        ordered_blocks.extend(tables_by_page.get(page_number, []))
    blocks = ordered_blocks
    _check_size(blocks)
    return (
        blocks,
        list(dict.fromkeys(warnings)),
        sorted({block.language for block in blocks}) or languages,
    )
