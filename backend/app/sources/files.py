"""Bounded, read-only ingestion and profiling for CSV and Excel sources."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import re
import zipfile
from collections.abc import Iterator, Sequence
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import PurePosixPath
from typing import Any, cast
from uuid import uuid4

import openpyxl
import olefile  # type: ignore[import-untyped]
import xlrd  # type: ignore[import-untyped]
from charset_normalizer import from_bytes
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db.models import Dataset, Source, Workspace
from app.storage.filesystem import Storage

_PROFILE_SCHEMA = "file-profile-v1"
_SAMPLE_ROWS = 10
_EXAMPLE_VALUES = 5
_TYPE_INFERENCE_ROWS = 10_000
_MAX_COLUMNS = 1_000
_MAX_WORKBOOK_MEMBERS = 4_096
_MAX_WORKBOOK_EXPANDED = 100 * 1024 * 1024
_MAX_WORKBOOK_ENTRY = 32 * 1024 * 1024
_MAX_WORKBOOK_CELLS = 1_000_000
_MAX_COMPRESSION_RATIO = 100
_MAX_PROFILE_CELL_CHARS = 512
_MAX_EXAMPLE_CHARS = 256
_MAX_PROFILE_SAMPLE_BYTES = 32 * 1024
_MAX_COLUMN_NAME_CHARS = 128
_MAX_ROW_CELL_CHARS = 4_096
_MAX_ROW_RESPONSE_BYTES = 1024 * 1024
_INTEGER = re.compile(r"^[+-]?(?:0|[1-9]\d*)$")
_ZERO_PADDED = re.compile(r"^[+-]?0\d+$")
_DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y")
csv.field_size_limit(25 * 1024 * 1024)


class FileIngestionError(ValueError):
    """Safe, user-facing rejection of an unsupported or malformed file."""

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code


class _ColumnStats:
    def __init__(
        self, name: str, position: int, example_count: int, example_limit: int
    ) -> None:
        self.name = name
        self.position = position
        self.missing_count = 0
        self.examples: list[str] = []
        self.observed_types: set[str] = set()
        self.inference_values = 0
        self.leading_zero = False
        self.ambiguous_dates = False
        self.example_truncated = False
        self.example_count = example_count
        self.example_limit = example_limit

    def add(self, value: object | None, row_number: int) -> None:
        if value is None or value == "":
            self.missing_count += 1
            return
        exported = _export_cell(value)
        text = exported if isinstance(exported, str) else str(exported)
        if text not in self.examples and len(self.examples) < self.example_count:
            if len(text) > self.example_limit:
                self.examples.append(text[: self.example_limit - 1] + "…")
                self.example_truncated = True
            else:
                self.examples.append(text)
        if row_number > _TYPE_INFERENCE_ROWS:
            return
        self.inference_values += 1
        inferred = _infer_value(value)
        self.observed_types.add(inferred)
        if isinstance(value, str) and _ZERO_PADDED.fullmatch(value.strip()):
            self.leading_zero = True
        if isinstance(value, str) and _looks_ambiguous_date(value.strip()):
            self.ambiguous_dates = True

    def finish(self, row_count: int) -> dict[str, Any]:
        types = self.observed_types
        uncertain = False
        if not types:
            logical = "unknown"
        elif self.leading_zero:
            logical, uncertain = "string", True
        elif types <= {"integer", "decimal"}:
            logical = "decimal" if "decimal" in types else "integer"
        elif len(types) == 1:
            logical = next(iter(types))
        else:
            logical, uncertain = "string", True
        duckdb = {
            "integer": "BIGINT",
            "decimal": "DECIMAL(38, 10)",
            "boolean": "BOOLEAN",
            "date": "DATE",
            "timestamp": "TIMESTAMP",
            "string": "VARCHAR",
            "unknown": "VARCHAR",
        }[logical]
        hints: dict[str, Any] = {}
        if self.leading_zero:
            hints["leading_zero_identifier"] = True
        if self.ambiguous_dates:
            hints["date_format_uncertain"] = True
        if self.name and re.search(r"(?i)(\binr\b|rupees?|₹)", self.name):
            hints["unit"] = "INR"
        elif self.name and re.search(
            r"(?i)(amount|income|salary|grant|total|price|cost)", self.name
        ):
            hints["currency_unknown"] = True
        if self.example_truncated:
            hints["example_values_truncated"] = True
        if uncertain:
            hints["type_inference_uncertain"] = True
        return {
            "name": self.name,
            "position": self.position,
            "type": logical,
            "duckdb_type": duckdb,
            "nullable": self.missing_count > 0 or row_count == 0,
            "missing_count": self.missing_count,
            "examples": self.examples,
            "hints": hints,
        }


def _infer_value(value: object) -> str:
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, datetime):
        return "timestamp"
    if isinstance(value, date):
        return "date"
    if isinstance(value, int):
        return "integer"
    if isinstance(value, (float, Decimal)):
        return "decimal"
    if not isinstance(value, str):
        return "string"
    text = value.strip()
    if text.lower() in {"true", "false"}:
        return "boolean"
    if _ZERO_PADDED.fullmatch(text):
        return "string"
    if _INTEGER.fullmatch(text):
        return "integer"
    if text:
        try:
            decimal = Decimal(text)
            if decimal.is_finite():
                return "decimal"
        except InvalidOperation:
            pass
        if "T" in text or " " in text:
            try:
                datetime.fromisoformat(text.replace("Z", "+00:00"))
                return "timestamp"
            except ValueError:
                pass
        if _parse_date(text) is not None:
            return "date"
    return "string"


def _parse_date(text: str) -> date | None:
    try:
        return date.fromisoformat(text)
    except ValueError:
        pass
    for fmt in _DATE_FORMATS[1:]:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _looks_ambiguous_date(text: str) -> bool:
    for fmt in ("%d/%m/%Y", "%m/%d/%Y"):
        try:
            datetime.strptime(text, fmt)
            return True
        except ValueError:
            continue
    return False


def _export_cell(value: object | None) -> object | None:
    if value is None or isinstance(value, (str, bool, int)):
        return value
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, float):
        if not math.isfinite(value):
            return None
        # Excel stores numeric values as IEEE-754 doubles. Return a decimal string
        # to avoid a second, silent float conversion in the API or JSON encoder.
        return str(Decimal(str(value)))
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def _cell_string(value: object | None) -> str:
    exported = _export_cell(value)
    return "" if exported is None else str(exported)


def _normalize_header(value: object | None, position: int, used: set[str]) -> str:
    base = str(value).strip() if value is not None else ""
    base = base or f"column_{position + 1}"
    base = base[:_MAX_COLUMN_NAME_CHARS]
    name = base
    suffix = 2
    while name.casefold() in used:
        name = f"{base}_{suffix}"
        suffix += 1
    used.add(name.casefold())
    return name


def _row_has_values(row: Sequence[object | None]) -> bool:
    return any(value is not None and value != "" for value in row)


def _detect_encoding(content: bytes) -> str:
    if content.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig"
    if content.startswith((b"\xff\xfe", b"\xfe\xff")):
        return "utf-16"
    match = from_bytes(content[:1_048_576]).best()
    if match is None or not match.encoding:
        raise FileIngestionError(
            "encoding_unknown", "CSV text encoding could not be detected."
        )
    try:
        content.decode(match.encoding, errors="strict")
    except (UnicodeError, LookupError) as error:
        raise FileIngestionError(
            "encoding_invalid", "CSV text could not be decoded safely."
        ) from error
    return match.encoding


def _detect_delimiter(text: str) -> str:
    sample = text[:65_536]
    if not sample.strip():
        raise FileIngestionError("csv_empty", "CSV file has no header or data rows.")
    try:
        return csv.Sniffer().sniff(sample, delimiters=",\t;|").delimiter
    except csv.Error:
        counts = {
            delimiter: sample.count(delimiter) for delimiter in (",", "\t", ";", "|")
        }
        candidate, count = max(counts.items(), key=lambda pair: pair[1])
        if count == 0:
            return ","  # Valid one-column CSV has no delimiter to sniff.
        return candidate


def _csv_rows(
    content: bytes, encoding: str, delimiter: str
) -> Iterator[list[object | None]]:
    try:
        text = content.decode(encoding, errors="strict")
        reader = csv.reader(
            io.StringIO(text, newline=""), delimiter=delimiter, strict=True
        )
        for row in reader:
            if row:
                yield list(row)
    except (UnicodeError, csv.Error) as error:
        raise FileIngestionError(
            "csv_invalid", "CSV content is malformed or contains invalid text."
        ) from error


def _profile_rows(
    identity: str,
    rows: Iterator[list[object | None]],
    *,
    encoding: str | None,
    delimiter: str | None,
    warnings: list[str] | None = None,
    header_note: str | None = None,
) -> dict[str, Any]:
    notes = list(dict.fromkeys(warnings or []))
    header: list[object | None] | None = None
    for candidate in rows:
        if _row_has_values(candidate):
            header = list(candidate)
            while header and (header[-1] is None or header[-1] == ""):
                header.pop()
            break
    if header is None:
        return {
            "identity": identity,
            "schema_version": _PROFILE_SCHEMA,
            "designation": "original",
            "storage_key": None,
            "lineage": [],
            "details": {
                "sheet_name": identity if identity != "data" else None,
                "header_row": None,
                "row_count": 0,
                "row_count_exact": True,
                "encoding": encoding,
                "delimiter": delimiter,
                "columns": [],
                "sample": [],
                "sample_truncated_cell_count": 0,
                "warnings": notes + ["empty_dataset"],
            },
        }
    if len(header) > _MAX_COLUMNS:
        raise FileIngestionError(
            "too_many_columns", f"A dataset may not exceed {_MAX_COLUMNS} columns."
        )
    used: set[str] = set()
    truncated_header_count = sum(
        1
        for value in header
        if value is not None and len(str(value).strip()) > _MAX_COLUMN_NAME_CHARS
    )
    names = [_normalize_header(item, index, used) for index, item in enumerate(header)]
    example_count = max(
        1,
        min(
            _EXAMPLE_VALUES,
            _MAX_PROFILE_SAMPLE_BYTES // max(1, 2 * len(names) * _MAX_EXAMPLE_CHARS),
        ),
    )
    example_limit = max(
        2,
        min(
            _MAX_EXAMPLE_CHARS,
            _MAX_PROFILE_SAMPLE_BYTES // max(1, 4 * len(names) * example_count),
        ),
    )
    stats = [
        _ColumnStats(name, index, example_count, example_limit)
        for index, name in enumerate(names)
    ]
    sample: list[dict[str, object | None]] = []
    truncated_sample_cells = 0
    sample_size_limited = False
    profile_cell_limit = max(
        1,
        min(
            _MAX_PROFILE_CELL_CHARS,
            _MAX_PROFILE_SAMPLE_BYTES // max(1, 4 * len(names) * _SAMPLE_ROWS),
        ),
    )
    row_count = 0
    for row in rows:
        if not _row_has_values(row):
            continue
        if len(row) > len(names):
            extra_values = row[len(names) :]
            if _row_has_values(extra_values):
                if len(row) > _MAX_COLUMNS:
                    raise FileIngestionError(
                        "too_many_columns",
                        f"A dataset may not exceed {_MAX_COLUMNS} columns.",
                    )
                while len(names) < len(row):
                    index = len(names)
                    name = _normalize_header(None, index, used)
                    names.append(name)
                    stats.append(
                        _ColumnStats(name, index, example_count, example_limit)
                    )
        row_count += 1
        normalized = list(row[: len(names)]) + [None] * max(0, len(names) - len(row))
        exported = [_export_cell(value) for value in normalized]
        if len(sample) < _SAMPLE_ROWS:
            bounded: list[object | None] = []
            row_truncated_count = 0
            for value in exported:
                if isinstance(value, str) and len(value) > profile_cell_limit:
                    value = value[: profile_cell_limit - 1] + "…"
                    row_truncated_count += 1
                bounded.append(value)
            record = dict(zip(names, bounded, strict=True))
            proposed_sample = sample + [record]
            if (
                len(
                    json.dumps(
                        proposed_sample, ensure_ascii=False, separators=(",", ":")
                    ).encode("utf-8")
                )
                <= _MAX_PROFILE_SAMPLE_BYTES
            ):
                sample.append(record)
                truncated_sample_cells += row_truncated_count
            else:
                sample_size_limited = True
        for column, value in zip(stats, normalized, strict=True):
            column.add(value, row_count)
    if row_count > _TYPE_INFERENCE_ROWS:
        notes.append("type_inference_sampled_first_10000_rows")
    if header_note:
        notes.append(header_note)
    if sample_size_limited:
        notes.append("sample_size_limit_reached")
    if truncated_header_count:
        notes.append("column_names_truncated")
    details = {
        "sheet_name": identity if identity != "data" else None,
        "header_row": 1,
        "row_count": row_count,
        "row_count_exact": True,
        "encoding": encoding,
        "delimiter": delimiter,
        "columns": [column.finish(row_count) for column in stats],
        "sample": sample,
        "sample_truncated_cell_count": truncated_sample_cells,
        "sample_omitted_row_count": max(0, row_count - len(sample)),
        "column_names_truncated_count": truncated_header_count,
        "warnings": list(dict.fromkeys(notes)),
    }
    return {
        "identity": identity,
        "schema_version": _PROFILE_SCHEMA,
        "designation": "original",
        "storage_key": None,
        "lineage": [],
        "details": details,
    }


def _profile_csv(content: bytes) -> list[dict[str, Any]]:
    encoding = _detect_encoding(content)
    try:
        decoded = content.decode(encoding, errors="strict")
    except (UnicodeError, LookupError) as error:
        raise FileIngestionError(
            "encoding_invalid", "CSV text could not be decoded safely."
        ) from error
    delimiter = _detect_delimiter(decoded)
    profile = _profile_rows(
        "data",
        _csv_rows(content, encoding, delimiter),
        encoding=encoding,
        delimiter=delimiter,
    )
    return [profile]


def _check_xlsx_archive(content: bytes) -> tuple[list[str], list[str]]:
    if not zipfile.is_zipfile(io.BytesIO(content)):
        raise FileIngestionError(
            "xlsx_invalid", "XLSX file is not a valid OOXML workbook."
        )
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
            if not members or len(members) > _MAX_WORKBOOK_MEMBERS:
                raise FileIngestionError(
                    "xlsx_archive_limit", "Workbook has too many archive entries."
                )
            total = 0
            cell_count = 0
            warnings: list[str] = []
            member_names: set[str] = set()
            for member in members:
                # ZIP readers can disagree about which copy of a repeated member
                # wins. Reject ambiguity before openpyxl sees the archive.
                if member.filename in member_names:
                    raise FileIngestionError(
                        "xlsx_unsafe_archive",
                        "Workbook contains duplicate archive entries.",
                    )
                member_names.add(member.filename)
                path = PurePosixPath(member.filename)
                raw_parts = member.filename.rstrip("/").split("/")
                if (
                    path.is_absolute()
                    or ".." in path.parts
                    or any(part in {"", "."} for part in raw_parts)
                    or "\\" in member.filename
                    or "\x00" in member.filename
                ):
                    raise FileIngestionError(
                        "xlsx_unsafe_archive",
                        "Workbook contains an unsafe archive path.",
                    )
                if member.file_size > _MAX_WORKBOOK_ENTRY:
                    raise FileIngestionError(
                        "xlsx_archive_limit", "Workbook contains an oversized entry."
                    )
                total += member.file_size
                if total > _MAX_WORKBOOK_EXPANDED:
                    raise FileIngestionError(
                        "xlsx_archive_limit",
                        "Workbook expands beyond the safe size limit.",
                    )
                if member.file_size and (
                    member.compress_size == 0
                    or member.file_size / member.compress_size > _MAX_COMPRESSION_RATIO
                ):
                    raise FileIngestionError(
                        "xlsx_compression_ratio",
                        "Workbook compression ratio exceeds the safe limit.",
                    )
                lower = member.filename.casefold()
                if any(
                    token in lower
                    for token in (
                        "vbaproject",
                        "activex",
                        "ctrlprops",
                        "macrosheets",
                        "dialogsheets",
                    )
                ):
                    raise FileIngestionError(
                        "macros_rejected",
                        "Macro-enabled or active-content workbooks are not accepted.",
                    )
                if "externallinks/" in lower:
                    warnings.append("external_links_ignored")
                if lower.endswith(".xml") or lower.endswith(".rels"):
                    raw = archive.read(member)
                    upper = raw.upper()
                    if b'TARGETMODE="EXTERNAL"' in upper:
                        warnings.append("external_links_ignored")
                    if b"<!DOCTYPE" in upper or b"<!ENTITY" in upper:
                        raise FileIngestionError(
                            "xml_entities_rejected",
                            "Workbook XML may not declare entities or doctypes.",
                        )
                    if lower.startswith("xl/worksheets/") and lower.endswith(".xml"):
                        cell_count += len(re.findall(rb"<c(?:\s|>)", raw))
                        if cell_count > _MAX_WORKBOOK_CELLS:
                            raise FileIngestionError(
                                "xlsx_cell_limit",
                                "Workbook contains too many populated cells.",
                            )
            return [member.filename for member in members], list(
                dict.fromkeys(warnings)
            )
    except zipfile.BadZipFile as error:
        raise FileIngestionError(
            "xlsx_invalid", "XLSX archive could not be read safely."
        ) from error


def _xlsx_value(value_cell: Any, formula_cell: Any) -> object | None:
    value = value_cell.value
    if value is None:
        return None
    number_format = getattr(value_cell, "number_format", "")
    if (
        isinstance(value, (int, float))
        and not isinstance(value, bool)
        and re.fullmatch(r"0{2,}", number_format or "")
    ):
        return f"{int(value):0{len(number_format)}d}"
    return cast(object, value)


def _xlsx_sheet_rows(
    formula_sheet: Any, value_sheet: Any, formula_stats: list[int]
) -> Iterator[list[object | None]]:
    formula_rows = formula_sheet.iter_rows(values_only=False)
    value_rows = value_sheet.iter_rows(values_only=False)
    for formulas, values in zip(formula_rows, value_rows):
        row: list[object | None] = []
        for formula_cell, value_cell in zip(formulas, values):
            if formula_cell.data_type == "f":
                formula_stats[1] += 1
                if value_cell.value is None:
                    formula_stats[0] += 1
                row.append(_xlsx_value(value_cell, formula_cell))
            else:
                row.append(_xlsx_value(value_cell, formula_cell))
        yield row


def _profile_xlsx(content: bytes) -> list[dict[str, Any]]:
    _, archive_warnings = _check_xlsx_archive(content)
    try:
        formulas_book = openpyxl.load_workbook(
            io.BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
        values_book = openpyxl.load_workbook(
            io.BytesIO(content), read_only=True, data_only=True, keep_links=False
        )
    except Exception as error:
        raise FileIngestionError(
            "xlsx_invalid", "XLSX workbook could not be parsed."
        ) from error
    profiles: list[dict[str, Any]] = []
    try:
        for formulas_sheet, values_sheet in zip(
            formulas_book.worksheets, values_book.worksheets, strict=True
        ):
            formula_stats = [0, 0]
            sheet_warnings = list(archive_warnings)
            if formulas_sheet.sheet_state != "visible":
                sheet_warnings.append("hidden_sheet_included")
            rows = _xlsx_sheet_rows(formulas_sheet, values_sheet, formula_stats)
            profile = _profile_rows(
                formulas_sheet.title,
                rows,
                encoding=None,
                delimiter=None,
                warnings=sheet_warnings,
            )
            if formula_stats[0]:
                profile["details"]["warnings"].append("formula_cached_value_missing")
                profile["details"]["formula_cache_missing_count"] = formula_stats[0]
            elif formula_stats[1]:
                profile["details"]["warnings"].append(
                    "formula_cached_values_used_not_recalculated"
                )
            profiles.append(profile)
    finally:
        formulas_book.close()
        values_book.close()
    if not profiles:
        raise FileIngestionError("xlsx_empty", "Workbook contains no worksheets.")
    return profiles


def _xls_cell(book: xlrd.book.Book, cell: xlrd.sheet.Cell) -> object | None:
    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
        return None
    if cell.ctype == xlrd.XL_CELL_DATE:
        try:
            return cast(
                object, xlrd.xldate.xldate_as_datetime(cell.value, book.datemode)
            )
        except (ValueError, OverflowError):
            return cast(object, cell.value)
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    if cell.ctype == xlrd.XL_CELL_NUMBER:
        format_item = book.format_map.get(book.xf_list[cell.xf_index].format_key)
        number_format = format_item.format_str if format_item else ""
        if re.fullmatch(r"0{2,}", number_format or ""):
            return f"{int(cell.value):0{len(number_format)}d}"
        if float(cell.value).is_integer():
            return int(cell.value)
        return Decimal(str(cell.value))
    return cast(object, cell.value)


def _xls_sheet_rows(
    book: xlrd.book.Book, sheet: xlrd.sheet.Sheet
) -> Iterator[list[object | None]]:
    for row_number in range(sheet.nrows):
        yield [
            _xls_cell(book, sheet.cell(row_number, col)) for col in range(sheet.ncols)
        ]


def _profile_xls(content: bytes) -> list[dict[str, Any]]:
    try:
        workbook = xlrd.open_workbook(
            file_contents=content, on_demand=True, formatting_info=True
        )
    except Exception as error:
        raise FileIngestionError(
            "xls_invalid", "Legacy XLS workbook could not be parsed."
        ) from error
    profiles: list[dict[str, Any]] = []
    try:
        warnings = [
            "legacy_xls_formulas_not_recalculated",
            "legacy_xls_workbook_loaded_in_memory",
        ]
        for sheet in workbook.sheets():
            profiles.append(
                _profile_rows(
                    sheet.name,
                    _xls_sheet_rows(workbook, sheet),
                    encoding=None,
                    delimiter=None,
                    warnings=warnings,
                )
            )
    finally:
        workbook.release_resources()
    if not profiles:
        raise FileIngestionError("xls_empty", "Workbook contains no worksheets.")
    return profiles


def _extension(filename: str) -> str:
    basename = PurePosixPath(filename.replace("\\", "/")).name
    if not basename or basename in {".", ".."} or basename.startswith("."):
        raise FileIngestionError(
            "filename_invalid", "File name must include a supported extension."
        )
    suffix = PurePosixPath(basename).suffix.lower()
    if suffix == ".xlsm":
        raise FileIngestionError(
            "macros_rejected", "Macro-enabled Excel workbooks are not accepted."
        )
    extension = suffix.removeprefix(".")
    if extension not in {"csv", "xlsx", "xls", "json", "parquet"}:
        raise FileIngestionError(
            "file_type_unsupported",
            "Only CSV, XLSX, legacy XLS, flat JSON, and Parquet files are supported.",
        )
    return extension


def _validate_content(extension: str, content: bytes) -> None:
    if not content:
        raise FileIngestionError("file_empty", "Uploaded file is empty.")
    if extension == "xlsx":
        _check_xlsx_archive(content)
    elif extension == "xls":
        if not content.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
            raise FileIngestionError(
                "xls_invalid", "Legacy XLS file signature is invalid."
            )
        try:
            with olefile.OleFileIO(io.BytesIO(content)) as compound_file:
                streams = compound_file.listdir(streams=True, storages=True)
                has_macros = any(
                    component.casefold() == "vba"
                    or "_vba_project_cur" in component.casefold()
                    or "macros" in component.casefold()
                    for path in streams
                    for component in path
                )
        except Exception as error:
            raise FileIngestionError(
                "xls_invalid", "Legacy XLS compound file could not be read safely."
            ) from error
        if has_macros:
            raise FileIngestionError(
                "macros_rejected",
                "Macro-enabled or active-content workbooks are not accepted.",
            )
    elif extension == "csv":
        if content.startswith((b"PK\x03\x04", bytes.fromhex("D0CF11E0A1B11AE1"))):
            raise FileIngestionError(
                "file_type_mismatch", "Uploaded bytes do not match the CSV extension."
            )


def profile_upload(filename: str, content: bytes) -> list[dict[str, Any]]:
    extension = _extension(filename)
    if extension != "xlsx":
        _validate_content(extension, content)
    if extension in {"json", "parquet"}:
        from app.sources.records import records_csv

        try:
            canonical = records_csv(extension, content)
        except Exception as error:
            raise FileIngestionError(
                "records_invalid",
                "Flat record file could not be decoded within safe bounds.",
            ) from error
        return _profile_csv(canonical)
    if extension == "csv":
        return _profile_csv(content)
    if extension == "xlsx":
        return _profile_xlsx(content)
    return _profile_xls(content)


def _iter_dataset_rows(
    filename: str, content: bytes, identity: str
) -> Iterator[list[object | None]]:
    extension = _extension(filename)
    _validate_content(extension, content)
    if extension in {"json", "parquet"}:
        from app.sources.records import records_csv

        try:
            content = records_csv(extension, content)
        except Exception as error:
            raise FileIngestionError(
                "records_invalid",
                "Flat record file could not be decoded within safe bounds.",
            ) from error
        extension = "csv"
    if extension == "csv":
        encoding = _detect_encoding(content)
        delimiter = _detect_delimiter(content.decode(encoding, errors="strict"))
        iterator = _csv_rows(content, encoding, delimiter)
        yield from _after_header(iterator)
        return
    if extension == "xlsx":
        formulas_book = openpyxl.load_workbook(
            io.BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
        values_book = openpyxl.load_workbook(
            io.BytesIO(content), read_only=True, data_only=True, keep_links=False
        )
        try:
            if identity not in formulas_book.sheetnames:
                raise FileIngestionError(
                    "dataset_sheet_not_found",
                    "Dataset sheet no longer exists in the original workbook.",
                )
            formula_sheet = formulas_book[identity]
            value_sheet = values_book[identity]
            rows = _xlsx_sheet_rows(formula_sheet, value_sheet, [0, 0])
            yield from _after_header(rows)
        finally:
            formulas_book.close()
            values_book.close()
        return
    book = xlrd.open_workbook(
        file_contents=content, on_demand=True, formatting_info=True
    )
    try:
        try:
            sheet = book.sheet_by_name(identity)
        except xlrd.biffh.XLRDError as error:
            raise FileIngestionError(
                "dataset_sheet_not_found",
                "Dataset sheet no longer exists in the original workbook.",
            ) from error
        rows = _xls_sheet_rows(book, sheet)
        yield from _after_header(rows)
    finally:
        book.release_resources()


def _header_for_dataset(dataset: Dataset) -> list[str]:
    profile_details = dataset.details.get("details", dataset.details)
    if not isinstance(profile_details, dict):
        raise FileIngestionError(
            "dataset_profile_invalid", "Stored dataset profile is invalid."
        )
    columns = profile_details.get("columns", [])
    if not isinstance(columns, list):
        raise FileIngestionError(
            "dataset_profile_invalid", "Stored dataset profile is invalid."
        )
    names = [
        str(item["name"])
        for item in columns
        if isinstance(item, dict) and "name" in item
    ]
    return names


def _after_header(
    rows: Iterator[list[object | None]],
) -> Iterator[list[object | None]]:
    for row in rows:
        if _row_has_values(row):
            break
    else:
        return
    yield from rows


def working_csv(
    source: Source, dataset: Dataset, storage: Storage, max_bytes: int
) -> bytes:
    """Return a bounded canonical UTF-8 CSV copy, preserving cell text and row order."""
    if max_bytes <= 0 or max_bytes > 100 * 1024 * 1024:
        raise ValueError("max_bytes must be between 1 and 100 MiB")
    if source.state != "ready" or source.storage_key is None:
        raise FileIngestionError(
            "source_unavailable", "Source is not ready for analysis."
        )
    if dataset.source_id != source.id or dataset.source_version != source.version:
        raise FileIngestionError(
            "dataset_source_mismatch", "Dataset does not match this source version."
        )
    if dataset.designation != "original":
        if dataset.storage_key is None:
            raise FileIngestionError(
                "dataset_storage_missing", "Derived dataset has no stored content."
            )
        return storage.read(dataset.storage_key, max_bytes)
    original = storage.read(source.storage_key, max_bytes)
    if (
        source.content_hash
        and hashlib.sha256(original).hexdigest() != source.content_hash
    ):
        raise FileIngestionError(
            "source_integrity_failed", "Original source hash verification failed."
        )
    output = io.BytesIO()
    stream = io.TextIOWrapper(output, encoding="utf-8", newline="", write_through=True)
    writer = csv.writer(stream, lineterminator="\n")
    columns = _header_for_dataset(dataset)
    writer.writerow(columns)
    if output.tell() > max_bytes:
        raise FileIngestionError(
            "prepared_size_limit", "Prepared dataset exceeds the size limit."
        )
    for row in _iter_dataset_rows(source.display_name, original, dataset.identity):
        values = list(row[: len(columns)]) + [None] * max(0, len(columns) - len(row))
        writer.writerow([_cell_string(value) for value in values])
        if output.tell() > max_bytes:
            raise FileIngestionError(
                "prepared_size_limit", "Prepared dataset exceeds the size limit."
            )
    stream.flush()
    return output.getvalue()


def ingest_file(
    session: Session,
    storage: Storage,
    workspace_id: str,
    filename: str,
    content: bytes,
    *,
    max_bytes: int,
) -> tuple[Source, list[Dataset]]:
    """Persist immutable original, bounded profiles, and ready source atomically in SQL."""
    if len(content) > max_bytes:
        raise FileIngestionError(
            "upload_too_large", "Uploaded file exceeds the configured size limit."
        )
    extension = _extension(filename)
    safe_filename = PurePosixPath(filename.replace("\\", "/")).name
    workspace = session.scalar(
        select(Workspace).where(Workspace.id == workspace_id).with_for_update()
    )
    if workspace is None:
        raise FileIngestionError("workspace_not_found", "Workspace not found.")
    source_id = str(uuid4())
    original_key = f"originals/{workspace_id}/{source_id}/original.{extension}"
    stored = storage.put(original_key, content)
    source = Source(
        id=source_id,
        workspace_id=workspace_id,
        kind=extension,
        version=1,
        display_name=safe_filename,
        state="processing",
        storage_key=stored.key,
        content_hash=stored.sha256,
        schema_version=_PROFILE_SCHEMA,
        details={"byte_size": stored.byte_size, "sha256": stored.sha256},
    )
    session.add(source)
    session.flush()
    try:
        profiles = profile_upload(safe_filename, content)
        datasets = [
            Dataset(
                source_id=source.id,
                source_version=source.version,
                identity=profile["identity"],
                schema_version=profile["schema_version"],
                details=profile["details"],
                storage_key=None,
                designation="original",
                lineage=[f"source:{source.id}@v{source.version}"],
            )
            for profile in profiles
        ]
        session.add_all(datasets)
        source.state = "ready"
        source.details = {
            "byte_size": stored.byte_size,
            "sha256": stored.sha256,
            "dataset_count": len(datasets),
        }
    except FileIngestionError as error:
        datasets = []
        source.state = "failed"
        source.details = {
            "byte_size": stored.byte_size,
            "sha256": stored.sha256,
            "error": {"code": error.code, "message": str(error)},
        }
    session.commit()
    session.refresh(source)
    for dataset in datasets:
        session.refresh(dataset)
    return source, datasets


def get_source_datasets(
    session: Session, source_id: str
) -> tuple[Source | None, list[Dataset]]:
    source = session.get(Source, source_id)
    if source is None:
        return None, []
    datasets = list(
        session.scalars(
            select(Dataset)
            .where(
                Dataset.source_id == source.id, Dataset.source_version == source.version
            )
            .order_by(Dataset.identity)
        ).all()
    )
    return source, datasets


def get_dataset_profile(
    session: Session, dataset_id: str
) -> tuple[Dataset | None, Source | None]:
    dataset = session.get(Dataset, dataset_id)
    if dataset is None:
        return None, None
    source = session.get(Source, dataset.source_id)
    if (
        source is None
        or source.version != dataset.source_version
        or source.state != "ready"
    ):
        return dataset, None
    return dataset, source


def get_dataset_rows(
    session: Session,
    storage: Storage,
    dataset_id: str,
    *,
    offset: int = 0,
    limit: int = 100,
    max_bytes: int,
    max_response_bytes: int = 65_536,
) -> dict[str, Any] | None:
    if offset < 0 or limit < 1 or limit > 500:
        raise ValueError(
            "offset must be nonnegative and limit must be between 1 and 500"
        )
    if max_response_bytes < 1_024 or max_response_bytes > _MAX_ROW_RESPONSE_BYTES:
        raise ValueError("max_response_bytes must be between 1024 and 1 MiB")
    dataset = session.get(Dataset, dataset_id)
    if dataset is None:
        return None
    source = session.get(Source, dataset.source_id)
    if (
        source is None
        or source.state != "ready"
        or source.version != dataset.source_version
    ):
        raise FileIngestionError("source_unavailable", "Dataset source is not ready.")
    if source.kind not in {"csv", "xlsx", "xls", "json", "parquet"}:
        raise FileIngestionError(
            "dataset_kind_unsupported",
            "Dataset rows are not stored as an uploaded file.",
        )
    if dataset.designation != "original":
        if dataset.storage_key is None:
            raise FileIngestionError(
                "dataset_storage_missing", "Derived dataset has no stored content."
            )
        raw = storage.read(dataset.storage_key, max_bytes)
        rows = _csv_rows(raw, "utf-8-sig", ",")
        next(rows, None)
    else:
        if source.storage_key is None:
            raise FileIngestionError(
                "source_storage_missing", "Original source bytes are unavailable."
            )
        raw = storage.read(source.storage_key, max_bytes)
        rows = _iter_dataset_rows(source.display_name, raw, dataset.identity)
    columns = _header_for_dataset(dataset)
    selected: list[dict[str, object | None]] = []
    truncated_cell_count = 0
    found = 0
    row_cell_limit = min(
        _MAX_ROW_CELL_CHARS,
        max(1, max_response_bytes // max(1, 16 * len(columns))),
    )
    for row in rows:
        if not _row_has_values(row):
            continue
        if found < offset:
            found += 1
            continue
        if len(selected) >= limit:
            break
        values = list(row[: len(columns)]) + [None] * max(0, len(columns) - len(row))
        exported = [_export_cell(value) for value in values]
        bounded: list[object | None] = []
        row_truncated = 0
        for value in exported:
            if isinstance(value, str) and len(value) > row_cell_limit:
                value = value[: row_cell_limit - 1] + "…"
                row_truncated += 1
            bounded.append(value)
        row_record = dict(zip(columns, bounded, strict=True))
        candidate_rows = selected + [row_record]
        candidate_bytes = len(
            json.dumps(
                candidate_rows, ensure_ascii=False, separators=(",", ":")
            ).encode("utf-8")
        )
        if candidate_bytes > max_response_bytes:
            break
        selected.append(row_record)
        truncated_cell_count += row_truncated
        found += 1
    profile_details = dataset.details.get("details", dataset.details)
    row_count = int(profile_details.get("row_count", 0))
    return {
        "dataset_id": dataset.id,
        "offset": offset,
        "limit": limit,
        "total_rows": row_count,
        "rows": selected,
        "truncated": offset + len(selected) < row_count,
        "truncated_cells": truncated_cell_count > 0,
        "truncated_cell_count": truncated_cell_count,
    }
