from __future__ import annotations

import io
from pathlib import Path

import pytest
from docx import Document as WordDocument
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import func, select

from app.api.document_lifecycle import router as lifecycle_router
from app.api.document_tables import router as tables_router
from app.api.documents import router as documents_router
from app.api.evidence import router as evidence_router
from app.api.files import router as files_router
from app.config import Settings
from app.db.models import (
    Dataset,
    Document,
    DocumentChunk,
    Evidence,
    Run,
    Source,
    Thread,
)
from app.db.session import get_session
from app.retrieval.service import search
from app.sources.documents import process_document
from app.sources.files import get_dataset_rows
from app.storage.filesystem import FileStorage
from tests.integration.test_runtime import db_factory

pytestmark = pytest.mark.integration


def _sample_docx() -> bytes:
    word = WordDocument()
    word.add_heading("Awards", level=1)
    word.add_paragraph("Eligible applicants receive an annual grant award.")
    table = word.add_table(rows=3, cols=2)
    for row_index, cells in enumerate(
        [
            ["Applicant", "Amount"],
            ["A-01", "10000"],
            ["A-02", "12500"],
        ]
    ):
        for column_index, value in enumerate(cells):
            table.cell(row_index, column_index).text = value
    stream = io.BytesIO()
    word.save(stream)
    return stream.getvalue()


def _test_client(
    db_factory,
    settings: Settings,
    storage: FileStorage,
    monkeypatch: pytest.MonkeyPatch,
) -> TestClient:
    import app.api.document_tables as table_api
    import app.api.documents as document_api
    import app.api.files as files_api

    monkeypatch.setattr(document_api, "get_settings", lambda: settings)
    monkeypatch.setattr(document_api, "get_storage", lambda _: storage)
    monkeypatch.setattr(table_api, "get_settings", lambda: settings)
    monkeypatch.setattr(table_api, "get_storage", lambda _: storage)
    monkeypatch.setattr(files_api, "get_settings", lambda: settings)
    monkeypatch.setattr(files_api, "get_storage", lambda _: storage)
    app = FastAPI()
    app.include_router(documents_router)
    app.include_router(tables_router)
    app.include_router(files_router)
    app.include_router(lifecycle_router)
    app.include_router(evidence_router)

    def override_session():
        with db_factory() as session:
            yield session

    app.dependency_overrides[get_session] = override_session
    return TestClient(app)


def test_document_table_accept_resume_and_archive_lifecycle(
    db_factory, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    settings = Settings(_env_file=None, storage_root=tmp_path)
    storage = FileStorage(tmp_path)
    client = _test_client(db_factory, settings, storage, monkeypatch)
    with db_factory() as session, session.begin():
        from app.db.models import Workspace

        workspace = Workspace(label="ingestion lifecycle")
        session.add(workspace)
        session.flush()
        workspace_id = workspace.id

    uploaded = client.post(
        f"/api/workspaces/{workspace_id}/documents",
        files={
            "file": (
                "awards.docx",
                _sample_docx(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert uploaded.status_code == 202, uploaded.text
    document_id = uploaded.json()["document"]["id"]
    source_id = uploaded.json()["source"]["id"]
    processed = process_document(document_id, settings, db_factory)
    assert processed["state"] == "ready"

    candidates = client.get(f"/api/documents/{document_id}/tables")
    assert candidates.status_code == 200
    table = candidates.json()["tables"][0]
    assert table["table_id"] == "table-0"
    assert table["row_count"] == 2
    assert table["preview"][0]["cells"][0]["provenance"]["block_id"]

    accepted = client.post(f"/api/documents/{document_id}/tables/table-0/accept")
    assert accepted.status_code == 200, accepted.text
    accepted_body = accepted.json()
    selected = client.get(f"/api/sources/{accepted_body['source_id']}/datasets")
    assert selected.status_code == 200
    assert selected.json()[0]["id"] == accepted_body["dataset_id"]
    profile = client.get(f"/api/datasets/{accepted_body['dataset_id']}/profile")
    assert profile.status_code == 200
    assert profile.json()["details"]["columns"][1]["name"] == "Amount"
    selected_rows = client.get(f"/api/datasets/{accepted_body['dataset_id']}/rows")
    assert selected_rows.status_code == 200
    assert sum(int(row["Amount"]) for row in selected_rows.json()["rows"]) == 22_500
    with db_factory() as session:
        result = get_dataset_rows(
            session,
            storage,
            accepted_body["dataset_id"],
            offset=0,
            limit=10,
            max_bytes=settings.max_upload_bytes,
        )
        assert result is not None
        amounts = [int(row["Amount"]) for row in result["rows"]]
        assert sum(amounts) == 22_500
        dataset = session.get(Dataset, accepted_body["dataset_id"])
        assert dataset is not None
        assert any(
            item.startswith("document_block:") for item in dataset.details["lineage"]
        )
    repeated = client.post(f"/api/documents/{document_id}/tables/table-0/accept")
    assert repeated.json()["dataset_id"] == accepted_body["dataset_id"]

    # Simulate a worker interruption after extraction publication but before
    # document-ready publication, then confirm resume keeps the same chunk IDs.
    with db_factory() as session, session.begin():
        document = session.get(Document, document_id)
        assert document is not None
        document.state = "running"
        before_ids = list(
            session.scalars(
                select(DocumentChunk.id)
                .where(DocumentChunk.document_id == document_id)
                .order_by(DocumentChunk.ordinal)
            )
        )
    resumed = process_document(document_id, settings, db_factory)
    assert resumed["extraction_reused"] is True
    with db_factory() as session:
        after_ids = list(
            session.scalars(
                select(DocumentChunk.id)
                .where(DocumentChunk.document_id == document_id)
                .order_by(DocumentChunk.ordinal)
            )
        )
        assert after_ids == before_ids
        assert session.scalar(
            select(func.count())
            .select_from(DocumentChunk)
            .where(DocumentChunk.document_id == document_id)
        ) == len(before_ids)

        chunk = session.scalar(
            select(DocumentChunk).where(
                DocumentChunk.document_id == document_id,
                DocumentChunk.normalized_text.contains("eligible applicants"),
            )
        )
        assert chunk is not None
        thread = Thread(workspace_id=workspace_id, label="saved citation")
        session.add(thread)
        session.flush()
        run = Run(thread_id=thread.id, state="completed", config={})
        session.add(run)
        session.flush()
        evidence = Evidence(
            run_id=run.id,
            kind="document",
            source_ids=[source_id],
            details={
                "document_id": document_id,
                "source_version": 1,
                "chunk_id": chunk.id,
                "excerpt": chunk.text,
                "location": chunk.location,
            },
        )
        session.add(evidence)
        session.flush()
        evidence_id = evidence.id
        session.commit()
    with db_factory() as session:
        before_delete = search(
            session,
            "eligible applicants grant award",
            {source_id: 1},
            settings,
            mode="text",
        )
        assert before_delete["passages"]

    removed = client.delete(f"/api/documents/{document_id}")
    assert removed.status_code == 200
    with db_factory() as session:
        after_delete = search(
            session,
            "eligible applicants grant award",
            {source_id: 1},
            settings,
            mode="text",
        )
    assert after_delete["passages"] == []
    archived = client.get(f"/api/evidence/{evidence_id}")
    assert archived.status_code == 200
    assert archived.json()["source_state"] == "archived"
    assert archived.json()["excerpt"]
