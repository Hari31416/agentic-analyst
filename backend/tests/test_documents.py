from __future__ import annotations

import io
import zipfile
from pathlib import Path
from typing import Any

import pytest
from tokenizers import Tokenizer, models, pre_tokenizers
from docx import Document as WordDocument
from fastapi import FastAPI
from fastapi.testclient import TestClient
from reportlab.pdfgen import canvas  # type: ignore[import-untyped]
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.api.documents import router
from app.config import Settings
from app.db.models import (
    Base,
    Document,
    DocumentBlock,
    DocumentChunk,
    IndexGeneration,
    Job,
    Source,
    Workspace,
)
from app.db.session import get_session
from app.sources.documents import (
    CHUNK_BODY_TOKENS,
    CHUNK_OVERLAP_TOKENS,
    CHUNKER_VERSION,
    EXTRACTOR_VERSION,
    DocumentIngestionError,
    ExtractedBlock,
    TokenizerProfile,
    _make_chunks,
    _tokenizer_profile,
    _validate_docx_archive,
    extract_document,
    process_document,
    validate_document_upload,
)
from app.storage.filesystem import FileStorage


def _docx_bytes() -> bytes:
    document = WordDocument()
    document.add_heading("Eligibility", level=1)
    document.add_paragraph("The annual income is below the limit.")
    document.add_paragraph("आवेदक को सहायता मिल सकती है।")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Scheme"
    table.cell(0, 1).text = "Grant"
    table.cell(1, 0).text = "S1"
    table.cell(1, 1).text = "₹10,000"
    document.add_paragraph("This follows the table.")
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def _pdf_bytes(text: str | None = "The grant is for eligible families.") -> bytes:
    stream = io.BytesIO()
    page = canvas.Canvas(stream)
    page.setTitle("Evaluation")
    if text:
        page.drawString(72, 720, text)
    page.showPage()
    page.save()
    return stream.getvalue()


def _database() -> tuple[sessionmaker[Session], Session]:
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    chunks_table = Base.metadata.tables["document_chunks"]
    postgres_only = [
        index
        for index in chunks_table.indexes
        if index.name is not None and index.name.endswith("_fts")
    ]
    for index in postgres_only:
        chunks_table.indexes.remove(index)
    try:
        Base.metadata.create_all(engine)
    finally:
        chunks_table.indexes.update(postgres_only)
    sessions = sessionmaker(engine, expire_on_commit=False)
    return sessions, sessions()


def _settings(tmp_path: Path) -> Settings:
    return Settings(
        _env_file=None,
        storage_root=tmp_path,
        embedding_model=None,
        embedding_model_path=None,
        embedding_revision=None,
    )  # type: ignore[call-arg]


def test_upload_validation_matches_extension_and_rejects_unsafe_docx() -> None:
    assert validate_document_upload("criteria.pdf", _pdf_bytes(), 1024 * 1024) == "pdf"
    assert (
        validate_document_upload("criteria.docx", _docx_bytes(), 1024 * 1024) == "docx"
    )
    with pytest.raises(DocumentIngestionError, match="extension"):
        validate_document_upload("criteria.docx", _pdf_bytes(), 1024 * 1024)
    with pytest.raises(DocumentIngestionError, match="Supported documents"):
        validate_document_upload("data.csv", b"a,b\n1,2", 1024 * 1024)

    source = _docx_bytes()
    output = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(source)) as original,
        zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as hostile,
    ):
        for item in original.infolist():
            content = original.read(item.filename)
            if item.filename == "word/document.xml":
                content = b'<!DOCTYPE x [<!ENTITY e "expanded">]><x>&e;</x>'
            hostile.writestr(item.filename, content)
    with pytest.raises(DocumentIngestionError, match="prohibited XML"):
        _validate_docx_archive(output.getvalue())


def test_docx_extraction_preserves_unicode_headings_tables_and_order() -> None:
    blocks, warnings, languages = extract_document("criteria.docx", _docx_bytes())
    assert not warnings
    assert "hi-IN" in languages
    assert [block.kind for block in blocks] == [
        "heading",
        "paragraph",
        "paragraph",
        "table_row",
        "table_row",
        "paragraph",
    ]
    assert blocks[0].heading == "Eligibility"
    assert "आवेदक" in blocks[2].text
    assert blocks[4].location["table"] == 0
    assert blocks[4].location["row"] == 1
    assert blocks[4].location["cells"] == ["S1", "₹10,000"]
    assert "Scheme: S1" in blocks[4].text
    assert blocks[-1].text == "This follows the table."


def test_external_docx_relationships_are_ignored_and_reported() -> None:
    source = _docx_bytes()
    output = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(source)) as original,
        zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as updated,
    ):
        for item in original.infolist():
            content = original.read(item.filename)
            if item.filename == "word/_rels/document.xml.rels":
                content = content.replace(
                    b"</Relationships>",
                    b'<Relationship Id="rId999" Type="https://example.test/link" '
                    b'Target="https://example.test/" TargetMode="External"/>'
                    b"</Relationships>",
                )
            updated.writestr(item.filename, content)
    blocks, warnings, _ = extract_document("criteria.docx", output.getvalue())
    assert blocks
    assert warnings == ["External document links were not followed."]


def test_pdf_extracts_page_locations_and_reports_scanned_page_for_ocr() -> None:
    blocks, warnings, languages = extract_document("criteria.pdf", _pdf_bytes())
    assert len(blocks) == 1
    assert blocks[0].location["page"] == 1
    assert "grant" in blocks[0].text.lower()
    assert "en-IN" in languages
    assert not warnings

    blocks, warnings, languages = extract_document("scan.pdf", _pdf_bytes(None))
    assert blocks == []
    assert "ocr_needed" in warnings
    assert languages == []


def test_fallback_chunker_preserves_unicode_whitespace_and_byte_bound() -> None:
    from app.sources.documents import _block

    text = ("  देवनागरी शब्द  " * 500) + (" final exact phrase  " * 500)
    block = _block("paragraph", text, "Heading", {"paragraph": 2})
    chunks = _make_chunks([block])
    assert len(chunks) > 1
    assert all(chunk.token_count <= 480 for chunk in chunks)
    assert all(chunk.normalized_text == chunk.text.casefold() for chunk in chunks)
    assert all(chunk.heading == "Heading" for chunk in chunks)
    assert all(
        chunk.location
        == {
            "type": "paragraphs",
            "paragraph_start": 2,
            "paragraph_end": 2,
            "chunk_strategy": "structure",
        }
        for chunk in chunks
    )
    assert "final exact phrase" in chunks[-1].text
    overlap = chunks[0].text.encode("utf-8")[-48:].decode("utf-8", errors="ignore")
    assert chunks[1].text.startswith(overlap)
    assert "  देवनागरी शब्द  " in chunks[0].text


def test_tokenizer_chunker_groups_narrative_before_tables_and_keeps_anchors() -> None:
    tokenizer = Tokenizer(models.WordLevel({"[UNK]": 0}, unk_token="[UNK]"))
    tokenizer.pre_tokenizer = pre_tokenizers.WhitespaceSplit()
    blocks = [
        ExtractedBlock(
            "heading",
            "Eligibility",
            "Eligibility",
            {"paragraph": 0},
            "en-IN",
            ["Latin"],
        ),
        ExtractedBlock(
            "paragraph",
            "  S1 qualifies.  ",
            "Eligibility",
            {"paragraph": 1},
            "en-IN",
            ["Latin"],
        ),
        ExtractedBlock(
            "paragraph",
            "Annual income <= INR 200000.\n",
            "Eligibility",
            {"paragraph": 2},
            "en-IN",
            ["Latin"],
        ),
        ExtractedBlock(
            "table_row",
            "Scheme: S1 | Grant: INR 10000",
            "Eligibility",
            {"table": 0, "row": 1},
            "en-IN",
            ["Latin"],
        ),
        ExtractedBlock(
            "paragraph",
            "After the table.",
            "Eligibility",
            {"paragraph": 3},
            "en-IN",
            ["Latin"],
        ),
    ]
    chunks = _make_chunks(
        blocks, TokenizerProfile(tokenizer, "test-tokenizer", "a" * 64)
    )
    narrative = chunks[0]
    assert narrative.text == "  S1 qualifies.  \n\nAnnual income <= INR 200000.\n"
    assert narrative.block_ids == ["0", "1", "2"]
    assert narrative.location == {
        "type": "paragraphs",
        "paragraph_start": 1,
        "paragraph_end": 2,
        "chunk_strategy": "structure",
    }
    table = chunks[1]
    assert "Scheme: S1" in table.text
    assert table.block_ids == ["0", "3"]
    assert table.location == {
        "type": "table",
        "table": 0,
        "row_start": 1,
        "row_end": 1,
        "cell_rows": [],
        "chunk_strategy": "structure",
    }
    assert chunks[2].block_ids == ["0", "4"]

    long_block = ExtractedBlock(
        "paragraph",
        "rule " * 600 + "FINAL-TAIL",
        "Eligibility",
        {"paragraph": 4},
        "en-IN",
        ["Latin"],
    )
    windows = _make_chunks(
        [long_block], TokenizerProfile(tokenizer, "test-tokenizer", None)
    )
    assert len(windows) > 1
    assert all(
        CHUNK_BODY_TOKENS
        <= window.token_count
        <= CHUNK_BODY_TOKENS + CHUNK_OVERLAP_TOKENS
        for window in windows[:-1]
    )
    assert (
        windows[1].text.split()[:CHUNK_OVERLAP_TOKENS]
        == windows[0].text.split()[-CHUNK_OVERLAP_TOKENS:]
    )
    assert "FINAL-TAIL" in windows[-1].text

    section_blocks = [
        ExtractedBlock(
            "paragraph",
            "start " * 50,
            "Eligibility",
            {"paragraph": 1},
            "en-IN",
            ["Latin"],
        ),
        ExtractedBlock(
            "paragraph",
            "middle " * 500,
            "Eligibility",
            {"paragraph": 2},
            "en-IN",
            ["Latin"],
        ),
        ExtractedBlock(
            "paragraph", "tail", "Eligibility", {"paragraph": 3}, "en-IN", ["Latin"]
        ),
    ]
    section_chunks = _make_chunks(
        section_blocks, TokenizerProfile(tokenizer, "test-tokenizer", None)
    )
    assert section_chunks[0].location == {
        "type": "paragraphs",
        "paragraph_start": 1,
        "paragraph_end": 2,
        "chunk_strategy": "structure",
    }
    assert section_chunks[-1].location["paragraph_start"] == 2
    assert section_chunks[-1].location["paragraph_end"] == 3


def test_configured_tokenizer_requires_manifest_hash_and_revision(
    tmp_path: Path,
) -> None:
    tokenizer = Tokenizer(models.WordLevel({"[UNK]": 0}, unk_token="[UNK]"))
    tokenizer.pre_tokenizer = pre_tokenizers.WhitespaceSplit()
    tokenizer.enable_truncation(max_length=2)
    tokenizer_path = tmp_path / "tokenizer.json"
    tokenizer.save(str(tokenizer_path))
    import hashlib
    import json

    digest = hashlib.sha256(tokenizer_path.read_bytes()).hexdigest()
    (tmp_path / "analyst-model.json").write_text(
        json.dumps(
            {
                "model_id": "intfloat/multilingual-e5-small",
                "revision": "test-revision",
                "sha256": {"tokenizer.json": digest},
            }
        ),
        encoding="utf-8",
    )
    settings = Settings(
        _env_file=None,
        embedding_model="intfloat/multilingual-e5-small",
        embedding_model_path=tmp_path,
        embedding_revision="test-revision",
    )  # type: ignore[call-arg]
    profile = _tokenizer_profile(settings)
    assert profile.tokenizer is not None
    assert profile.method == "tokenizer-json-v2"
    assert profile.sha256 == digest
    assert (
        len(
            profile.tokenizer.encode("one two three four", add_special_tokens=False).ids
        )
        == 4
    )
    settings.embedding_revision = "wrong-revision"
    fallback = _tokenizer_profile(settings)
    assert fallback.tokenizer is None
    assert fallback.method == "utf8-byte-fallback-480-v2"


def _add_document(
    sessions: sessionmaker[Session], storage: FileStorage, content: bytes
) -> str:
    with sessions() as session:
        workspace = Workspace(label="Documents")
        session.add(workspace)
        session.flush()
        source = Source(
            workspace_id=workspace.id,
            kind="document",
            version=1,
            display_name="criteria.pdf",
            state="processing",
            content_hash="a" * 64,
            details={"media_type": "application/pdf"},
        )
        session.add(source)
        session.flush()
        stored = storage.put(
            f"originals/{workspace.id}/{source.id}/original.pdf", content
        )
        source.storage_key = stored.key
        document = Document(
            source_id=source.id,
            source_version=source.version,
            extractor_version="pdf-docx-text-v1",
            chunker_version=CHUNKER_VERSION,
            state="queued",
            stage="queued",
            progress=0,
            details={},
        )
        session.add(document)
        session.commit()
        return document.id


def test_process_document_persists_blocks_chunks_and_degraded_generation(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    sessions, session = _database()
    session.close()
    storage = FileStorage(tmp_path)
    document_id = _add_document(sessions, storage, _pdf_bytes())
    monkeypatch.setattr("app.storage.factory.get_storage", lambda settings: storage)
    from app.retrieval.embedding import EmbeddingUnavailable

    monkeypatch.setattr(
        "app.retrieval.service.get_embedding_adapter",
        lambda settings: (_ for _ in ()).throw(
            EmbeddingUnavailable("embedding_model_load_failed", "offline")
        ),
    )

    guarded_sessions: list[Session] = []
    result = process_document(
        document_id,
        _settings(tmp_path),
        sessions,
        lease_guard=lambda current: guarded_sessions.append(current),
    )
    assert result["stage"] == "index_degraded"
    assert len(guarded_sessions) >= 4
    with sessions() as check:
        document = check.get(Document, document_id)
        assert document is not None
        assert document.state == "ready"
        assert document.stage == "index_degraded"
        assert document.details["index_status"] == "unavailable"
        assert document.index_generation_id is not None
        generation = check.get(IndexGeneration, document.index_generation_id)
        assert generation is not None and generation.status == "degraded"
        blocks = list(
            check.scalars(
                select(DocumentBlock).where(DocumentBlock.document_id == document_id)
            )
        )
        chunks = list(
            check.scalars(
                select(DocumentChunk).where(DocumentChunk.document_id == document_id)
            )
        )
        assert len(blocks) == len(chunks) == 1
        assert chunks[0].block_ids == [blocks[0].id]
        assert chunks[0].token_count == len(chunks[0].text.encode("utf-8"))


def test_lost_lease_before_index_publication_rolls_back_generation(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    sessions, session = _database()
    session.close()
    storage = FileStorage(tmp_path)
    document_id = _add_document(sessions, storage, _pdf_bytes())
    monkeypatch.setattr("app.storage.factory.get_storage", lambda settings: storage)
    calls = 0

    def guard(_: Session) -> None:
        nonlocal calls
        calls += 1
        if calls >= 4:
            raise RuntimeError("lease lost")

    with pytest.raises(RuntimeError, match="lease lost"):
        process_document(document_id, _settings(tmp_path), sessions, lease_guard=guard)
    with sessions() as check:
        assert (
            check.scalar(
                select(func.count())
                .select_from(DocumentBlock)
                .where(DocumentBlock.document_id == document_id)
            )
            == 1
        )
        assert (
            check.scalar(
                select(func.count())
                .select_from(DocumentChunk)
                .where(DocumentChunk.document_id == document_id)
            )
            == 1
        )
        assert (
            check.scalar(
                select(func.count())
                .select_from(IndexGeneration)
                .where(IndexGeneration.document_id == document_id)
            )
            == 0
        )
        document = check.get(Document, document_id)
        assert document is not None and document.stage == "indexing"


def test_chunk_limit_is_persisted_as_terminal_document_failure(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    sessions, session = _database()
    session.close()
    storage = FileStorage(tmp_path)
    document_id = _add_document(sessions, storage, _pdf_bytes())
    monkeypatch.setattr("app.storage.factory.get_storage", lambda _: storage)
    blocks = [
        ExtractedBlock(
            kind="paragraph",
            text="criterion " * 400,
            heading=None,
            location={"paragraph": index},
            language="en-IN",
            scripts=["Latin"],
        )
        for index in range(300)
    ]
    monkeypatch.setattr(
        "app.sources.documents.extract_document",
        lambda filename, content, **options: (blocks, [], ["en-IN"]),
    )

    result = process_document(
        document_id,
        _settings(tmp_path).model_copy(update={"document_max_chunks": 512}),
        sessions,
    )
    assert result["state"] == "failed"
    assert result["error"] == "document_chunk_limit"
    with sessions() as check:
        document = check.get(Document, document_id)
        assert document is not None
        assert document.state == "failed"
        assert document.stage == "failed"
        assert document.progress == 100
        assert document.details["error"]["code"] == "document_chunk_limit"
        source = check.get(Source, document.source_id)
        assert source is not None and source.state == "failed"
        assert (
            check.scalar(
                select(func.count())
                .select_from(DocumentBlock)
                .where(DocumentBlock.document_id == document_id)
            )
            == 0
        )


def test_document_api_upload_queues_ingestion_and_pages_blocks(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    sessions, session = _database()
    workspace = Workspace(label="API")
    session.add(workspace)
    session.commit()
    workspace_id = workspace.id
    session.close()
    settings = _settings(tmp_path)
    storage = FileStorage(tmp_path)
    monkeypatch.setattr("app.api.documents.get_settings", lambda: settings)
    monkeypatch.setattr("app.api.documents.get_storage", lambda _: storage)

    app = FastAPI()
    app.include_router(router)

    def override_session():
        with sessions() as active:
            yield active

    app.dependency_overrides[get_session] = override_session
    client = TestClient(app)
    content = _pdf_bytes()
    response = client.post(
        f"/api/workspaces/{workspace_id}/documents",
        files={"file": ("criteria.pdf", content, "application/pdf")},
    )
    assert response.status_code == 202, response.text
    payload: dict[str, Any] = response.json()
    document_id = payload["document"]["id"]
    assert payload["document"]["state"] == "queued"
    assert payload["source"]["state"] == "processing"
    duplicate = client.post(
        f"/api/workspaces/{workspace_id}/documents",
        files={"file": ("criteria.pdf", content, "application/pdf")},
    )
    assert duplicate.status_code == 202
    assert duplicate.json()["source"]["id"] == payload["source"]["id"]
    assert duplicate.json()["document"]["id"] == document_id
    with sessions() as check:
        job = check.scalar(
            select(Job).where(Job.payload["document_id"].as_string() == document_id)
        )
        assert job is not None and job.kind == "ingest_document"
        assert check.scalar(select(func.count()).select_from(Source)) == 1
        assert check.scalar(select(func.count()).select_from(Job)) == 1

    listed = client.get(f"/api/workspaces/{workspace_id}/documents")
    assert listed.status_code == 200
    assert listed.json()[0]["id"] == document_id
    blocks = client.get(f"/api/documents/{document_id}/blocks?offset=0&limit=2")
    assert blocks.status_code == 200
    assert blocks.json()["total"] == 0
    assert blocks.json()["blocks"] == []


def test_block_api_returns_bounded_pagination(tmp_path: Path) -> None:
    sessions, session = _database()
    workspace = Workspace(label="Blocks")
    session.add(workspace)
    session.flush()
    source = Source(
        workspace_id=workspace.id,
        kind="document",
        version=1,
        display_name="sample.docx",
        state="ready",
        content_hash="b" * 64,
        details={},
    )
    session.add(source)
    session.flush()
    document = Document(
        source_id=source.id,
        source_version=1,
        extractor_version="v1",
        chunker_version="v1",
        state="ready",
        stage="indexed",
        progress=100,
        details={},
    )
    session.add(document)
    session.flush()
    for ordinal in range(3):
        session.add(
            DocumentBlock(
                document_id=document.id,
                ordinal=ordinal,
                kind="paragraph",
                text=f"Paragraph {ordinal}",
                heading=None,
                location={"paragraph": ordinal},
                language="en-IN",
                scripts=["Latin"],
            )
        )
    session.commit()
    document_id = document.id
    session.close()

    app = FastAPI()
    app.include_router(router)

    def override_session():
        with sessions() as active:
            yield active

    app.dependency_overrides[get_session] = override_session
    response = TestClient(app).get(
        f"/api/documents/{document_id}/blocks?offset=1&limit=1"
    )
    assert response.status_code == 200
    assert response.json()["total"] == 3
    assert response.json()["blocks"][0]["ordinal"] == 1
    assert response.json()["next_offset"] == 2


def _image_bytes(fmt: str = "PNG", size: tuple[int, int] = (64, 64)) -> bytes:
    from PIL import Image

    img = Image.new("RGB", size, color=(255, 0, 0))
    buf = io.BytesIO()
    img.save(buf, format=fmt)
    return buf.getvalue()


def test_validate_document_upload_image() -> None:
    png_data = _image_bytes("PNG")
    kind = validate_document_upload("photo.png", png_data, len(png_data) + 100)
    assert kind == "image"

    jpeg_data = _image_bytes("JPEG")
    kind = validate_document_upload("photo.jpg", jpeg_data, len(jpeg_data) + 100)
    assert kind == "image"

    with pytest.raises(DocumentIngestionError, match="image"):
        validate_document_upload("bad.png", b"not an image", 1000)


def test_extract_document_standalone_image() -> None:
    png_data = _image_bytes("PNG")
    blocks, warnings, languages = extract_document(
        "figure.png", png_data, ocr_enabled=False
    )
    assert len(blocks) == 1
    assert blocks[0].kind == "image"
    assert blocks[0].image_bytes == png_data
    assert blocks[0].location["type"] == "image"
    assert blocks[0].location["image_format"] == "png"
    assert blocks[0].location["image_width"] == 64
    assert blocks[0].location["image_height"] == 64


def test_process_document_standalone_image(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    sessions, session = _database()
    session.close()
    storage = FileStorage(tmp_path)
    monkeypatch.setattr("app.storage.factory.get_storage", lambda settings: storage)
    monkeypatch.setattr("app.api.documents.get_storage", lambda _: storage)

    png_data = _image_bytes("PNG")
    with sessions() as s:
        workspace = Workspace(label="Images")
        s.add(workspace)
        s.flush()
        source = Source(
            workspace_id=workspace.id,
            kind="document",
            version=1,
            display_name="chart.png",
            state="processing",
            content_hash="c" * 64,
            details={"media_type": "image/png"},
        )
        s.add(source)
        s.flush()
        stored = storage.put(
            f"originals/{workspace.id}/{source.id}/original.png", png_data
        )
        source.storage_key = stored.key
        document = Document(
            source_id=source.id,
            source_version=source.version,
            extractor_version=EXTRACTOR_VERSION,
            chunker_version=CHUNKER_VERSION,
            state="queued",
            stage="queued",
            progress=0,
            details={},
        )
        s.add(document)
        s.commit()
        doc_id = document.id
        ws_id = workspace.id

    result = process_document(
        doc_id,
        _settings(tmp_path),
        sessions,
    )
    assert result["state"] == "ready"

    with sessions() as s:
        blocks = s.scalars(
            select(DocumentBlock).where(DocumentBlock.document_id == doc_id)
        ).all()
        assert len(blocks) == 1
        assert blocks[0].kind == "image"
        assert blocks[0].location["type"] == "image"
        img_key = blocks[0].location["image_key"]
        assert img_key.startswith(f"derived/{ws_id}/{doc_id}/images/")
        assert storage.read(img_key) == png_data

        chunks = s.scalars(
            select(DocumentChunk).where(DocumentChunk.document_id == doc_id)
        ).all()
        assert len(chunks) == 1
        assert chunks[0].location["type"] == "image"
        assert chunks[0].location["image_key"] == img_key

    app = FastAPI()
    app.include_router(router)

    def override_session():
        with sessions() as active:
            yield active

    app.dependency_overrides[get_session] = override_session
    client = TestClient(app)

    block_resp = client.get(f"/api/documents/{doc_id}/blocks/{blocks[0].id}/image")
    assert block_resp.status_code == 200
    assert block_resp.headers["content-type"] == "image/jpeg"
    assert len(block_resp.content) > 0

    doc_resp = client.get(f"/api/documents/{doc_id}/image")
    assert doc_resp.status_code == 200
    assert doc_resp.headers["content-type"] == "image/jpeg"
    assert len(doc_resp.content) > 0


def test_docx_extracts_embedded_media_images() -> None:
    doc = WordDocument()
    doc.add_paragraph("Paragraph before image.")
    img_stream = io.BytesIO(_image_bytes("PNG", size=(50, 50)))
    doc.add_picture(img_stream, width=None, height=None)
    doc.add_paragraph("Paragraph after image.")
    docx_buf = io.BytesIO()
    doc.save(docx_buf)

    blocks, warnings, languages = extract_document(
        "with_img.docx", docx_buf.getvalue(), ocr_enabled=False
    )
    image_blocks = [b for b in blocks if b.kind == "image"]
    assert len(image_blocks) >= 1
    assert image_blocks[0].image_bytes is not None
    assert image_blocks[0].location["type"] == "image"


@pytest.mark.parametrize("scope", ["source", "document", "workspace", "retained"])
def test_image_cleanup_obeys_source_retention(scope):
    from uuid import UUID
    from app.api.resource_lifecycle import delete_workspace, remove_source
    from app.api.document_lifecycle import remove_document
    from app.db.models import Run, Thread

    sessions, session = _database()
    workspace = Workspace(label="images")
    session.add(workspace)
    session.flush()
    source = Source(
        workspace_id=workspace.id,
        kind="document",
        display_name="chart.png",
        state="ready",
        storage_key="originals/chart.png",
        details={},
    )
    session.add(source)
    session.flush()
    document = Document(
        source_id=source.id,
        source_version=1,
        extractor_version=EXTRACTOR_VERSION,
        chunker_version=CHUNKER_VERSION,
        state="ready",
        stage="ready",
        progress=100,
        details={},
    )
    session.add(document)
    session.flush()
    image_key = f"derived/{workspace.id}/{document.id}/images/chart.png"
    session.add(
        DocumentBlock(
            document_id=document.id,
            ordinal=0,
            kind="image",
            text="chart",
            heading=None,
            location={"image_key": image_key},
            language="en-IN",
            scripts=[],
        )
    )
    session.add(
        DocumentChunk(
            document_id=document.id,
            ordinal=0,
            chunker_version=CHUNKER_VERSION,
            text="chart",
            normalized_text="chart",
            heading=None,
            location={"image_key": image_key},
            language="en-IN",
            block_ids=[],
            token_count=1,
        )
    )
    if scope == "retained":
        thread = Thread(workspace_id=workspace.id, label="citation")
        session.add(thread)
        session.flush()
        session.add(
            Run(
                thread_id=thread.id,
                state="completed",
                selected_source_ids=[source.id],
                config={},
            )
        )
    session.commit()
    if scope == "workspace":
        delete_workspace(UUID(workspace.id), session)
    elif scope == "document":
        remove_document(UUID(document.id), session)
    else:
        result = remove_source(session, source.id)
        assert result["retention"] == (
            "archived_for_citations" if scope == "retained" else "purged"
        )
    queued = list(session.scalars(select(Job).where(Job.kind == "delete_storage")))
    assert sum(job.payload["storage_key"] == image_key for job in queued) == (
        0 if scope == "retained" else 1
    )
    session.close()
