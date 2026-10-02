from __future__ import annotations

import csv
import io
import zipfile
from uuid import uuid4

import openpyxl
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.api.files import router
from app.config import Settings
from app.db.models import Base, Dataset, Source, Workspace
from app.db.session import get_session
from app.sources.files import (
    FileIngestionError,
    get_dataset_rows,
    ingest_file,
    profile_upload,
    working_csv,
)
from app.storage.filesystem import FileStorage


def _xlsx_bytes(*, formula: bool = False) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["ID", "amount", "date", "nullable"])
    sheet.append(["00123", "25000.00", "2026-01-02", None])
    sheet.append(["00124", "0.10", "2026-01-03", "ok"])
    if formula:
        sheet["E1"] = "computed"
        sheet["E2"] = "=1+1"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _xls_bytes() -> bytes:
    xlwt = pytest.importorskip("xlwt")
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("Legacy")
    for col, value in enumerate(["account", "amount"]):
        sheet.write(0, col, value)
    sheet.write(1, 0, "00123")
    sheet.write(1, 1, 10.25)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _csv_bytes(text: str, encoding: str = "utf-8") -> bytes:
    return text.encode(encoding)


def test_csv_profile_preserves_decimal_and_leading_zero_text():
    profiles = profile_upload(
        "records.csv",
        _csv_bytes("account,amount,nullable\n00123,25000.00,\n00124,0.10,ok\n"),
    )
    details = profiles[0]["details"]
    assert details["row_count"] == 2
    assert details["delimiter"] == ","
    assert details["columns"][0]["type"] == "string"
    assert details["columns"][0]["hints"]["leading_zero_identifier"] is True
    assert details["columns"][1]["type"] == "decimal"
    assert details["sample"][0]["amount"] == "25000.00"
    assert details["columns"][2]["missing_count"] == 1
    assert details["columns"][2]["nullable"] is True


def test_single_column_csv_and_long_cells_are_bounded():
    one_column = profile_upload("single.csv", b"code\nA\nB\n")[0]
    assert one_column["details"]["delimiter"] == ","
    assert one_column["details"]["row_count"] == 2


def test_csv_detects_legacy_encoding_and_semicolon_delimiter():
    profile = profile_upload(
        "records.csv", _csv_bytes("name;amount\nJos\u00e9;12.50\n", "cp1252")
    )[0]
    assert profile["details"]["delimiter"] == ";"
    assert profile["details"]["row_count"] == 1
    assert profile["details"]["sample"][0]["name"] == "Jos\u00e9"


def test_xlsx_profiles_sheets_dates_and_cached_formula_warning():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Applicants"
    sheet.append(["ID", "amount", "due"])
    sheet.append(["001", 0.1, __import__("datetime").date(2026, 2, 3)])
    sheet["D1"] = "result"
    sheet["D2"] = "=1+1"
    output = io.BytesIO()
    workbook.save(output)
    profiles = profile_upload("data.xlsx", output.getvalue())
    assert profiles[0]["identity"] == "Applicants"
    details = profiles[0]["details"]
    assert details["row_count"] == 1
    assert details["sample"][0]["ID"] == "001"
    assert details["sample"][0]["due"] == "2026-02-03T00:00:00"
    assert "formula_cached_value_missing" in details["warnings"]


def test_xlsx_macro_payload_and_archive_limits_rejected():
    content = _xlsx_bytes()
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        members = [
            (item.filename, archive.read(item.filename)) for item in archive.infolist()
        ]
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        for name, payload in members:
            archive.writestr(name, payload)
        archive.writestr("xl/vbaProject.bin", b"macro")
    with pytest.raises(FileIngestionError, match="Macro-enabled"):
        profile_upload("data.xlsx", output.getvalue())
    with pytest.raises(FileIngestionError):
        profile_upload("data.xlsm", content)


def test_xls_profiles_legacy_workbook():
    profile = profile_upload("legacy.xls", _xls_bytes())[0]
    assert profile["identity"] == "Legacy"
    assert profile["details"]["sample"][0]["account"] == "00123"
    assert "legacy_xls_workbook_loaded_in_memory" in profile["details"]["warnings"]


def test_profile_rejects_mismatched_types_and_entity_xml():
    with pytest.raises(FileIngestionError, match="do not match"):
        profile_upload("wrong.csv", _xlsx_bytes())
    content = _xlsx_bytes()
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        members = [
            (item.filename, archive.read(item.filename)) for item in archive.infolist()
        ]
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        for name, payload in members:
            if name == "xl/workbook.xml":
                payload = b'<!DOCTYPE x [<!ENTITY a "x">]>' + payload
            archive.writestr(name, payload)
    with pytest.raises(FileIngestionError, match="entities"):
        profile_upload("bad.xlsx", output.getvalue())


@pytest.fixture
def file_api(tmp_path, monkeypatch):
    engine = create_engine(
        "sqlite://", poolclass=StaticPool, connect_args={"check_same_thread": False}
    )
    Base.metadata.create_all(engine)
    sessions = sessionmaker(engine, expire_on_commit=False)

    def dependency():
        with sessions() as session:
            yield session

    app = FastAPI()
    app.include_router(router)
    app.dependency_overrides[get_session] = dependency
    settings = Settings(_env_file=None, storage_root=tmp_path)
    monkeypatch.setattr("app.api.files.get_settings", lambda: settings)
    monkeypatch.setattr(
        "app.api.files.get_storage", lambda _settings: FileStorage(tmp_path)
    )
    with sessions() as session:
        workspace = Workspace(label="test")
        session.add(workspace)
        session.commit()
        workspace_id = workspace.id
    try:
        yield TestClient(app), sessions, workspace_id, FileStorage(tmp_path)
    finally:
        app.dependency_overrides.clear()
        engine.dispose()


def test_upload_api_persists_original_and_profiles_before_ready(file_api):
    client, sessions, workspace_id, storage = file_api
    response = client.post(
        f"/api/workspaces/{workspace_id}/sources/files",
        files={"file": ("awards.csv", b"name,amount\nA,10.00\nB,15.00\n", "text/csv")},
    )
    assert response.status_code == 201, response.text
    source_payload = response.json()
    assert source_payload["state"] == "ready"
    assert source_payload["kind"] == "csv"
    assert source_payload["datasets"][0]["details"]["row_count"] == 2
    with sessions() as session:
        source = session.get(Source, source_payload["id"])
        dataset = session.scalar(select(Dataset).where(Dataset.source_id == source.id))
        assert source.storage_key.startswith("originals/")
        assert storage.read(source.storage_key) == b"name,amount\nA,10.00\nB,15.00\n"
        assert dataset.designation == "original"
        assert dataset.storage_key is None
        assert source.state == "ready"


def test_dataset_rows_and_working_csv_stream_original_bytes(file_api):
    _, sessions, workspace_id, storage = file_api
    with sessions() as session:
        source, datasets = ingest_file(
            session,
            storage,
            workspace_id,
            "grants.csv",
            b"ID,amount\n001,10000.00\n002,15000.00\n003,5\n",
            max_bytes=1024,
        )
        page = get_dataset_rows(
            session, storage, datasets[0].id, offset=1, limit=1, max_bytes=1024
        )
        assert page["rows"] == [{"ID": "002", "amount": "15000.00"}]
        assert page["total_rows"] == 3
        prepared = working_csv(source, datasets[0], storage, 1024)
        assert prepared.startswith(b"ID,amount\n001,10000.00\n")


def test_large_sample_cells_report_truncation_and_bounded_rows(file_api):
    client, sessions, workspace_id, storage = file_api
    raw = b"id,payload\n1," + b"x" * 10_000 + b"\n"
    with sessions() as session:
        source, datasets = ingest_file(
            session,
            storage,
            workspace_id,
            "large.csv",
            raw,
            max_bytes=20_000,
        )
        details = datasets[0].details
        assert details["sample_truncated_cell_count"] == 1
        assert details["columns"][1]["hints"]["example_values_truncated"]
        page = get_dataset_rows(
            session,
            storage,
            datasets[0].id,
            offset=0,
            limit=1,
            max_bytes=20_000,
            max_response_bytes=4_096,
        )
        assert page["truncated_cells"] is True
        assert page["truncated_cell_count"] == 1
        assert len(page["rows"][0]["payload"]) <= 128


def test_xlsx_upload_creates_sheet_datasets_and_samples(file_api):
    client, _, workspace_id, _ = file_api
    response = client.post(
        f"/api/workspaces/{workspace_id}/sources/files",
        files={"file": ("records.xlsx", _xlsx_bytes(), "application/octet-stream")},
    )
    assert response.status_code == 201, response.text
    payload = response.json()
    assert payload["state"] == "ready"
    assert [item["identity"] for item in payload["datasets"]] == ["Data"]
    dataset_id = payload["datasets"][0]["id"]
    profile = client.get(f"/api/datasets/{dataset_id}/profile")
    rows = client.get(f"/api/datasets/{dataset_id}/rows?offset=1&limit=1")
    assert profile.status_code == rows.status_code == 200
    assert rows.json()["rows"][0]["ID"] == "00124"


def test_registered_derived_dataset_without_extension_is_sampled(file_api):
    _, sessions, workspace_id, storage = file_api
    content = b"account,amount\n001,10.00\n002,20.00\n"
    profile = profile_upload("result.csv", content)[0]
    with sessions() as session:
        source_id = str(uuid4())
        key = f"derived/{workspace_id}/{source_id}/prepared/result.csv"
        stored = storage.put(key, content)
        source = Source(
            id=source_id,
            workspace_id=workspace_id,
            kind="csv",
            version=1,
            display_name="Derived awards",
            state="ready",
            storage_key=key,
            content_hash=stored.sha256,
            details={"designation": "derived"},
        )
        session.add(source)
        session.flush()
        dataset = Dataset(
            source_id=source.id,
            source_version=source.version,
            identity="data",
            schema_version=profile["schema_version"],
            details=profile,
            storage_key=key,
            designation="derived",
            lineage=["artifact:derived"],
        )
        session.add(dataset)
        session.commit()
        page = get_dataset_rows(
            session, storage, dataset.id, offset=0, limit=1, max_bytes=1024
        )
        assert page["rows"] == [{"account": "001", "amount": "10.00"}]
        assert working_csv(source, dataset, storage, 1024) == content


def test_source_metadata_patch_preserves_file_profile_and_workspace_scope(file_api):
    client, sessions, workspace_id, _ = file_api
    uploaded = client.post(
        f"/api/workspaces/{workspace_id}/sources/files",
        files={"file": ("source.csv", b"id,amount\n1,20.00\n", "text/csv")},
    )
    assert uploaded.status_code == 201
    source_id = uploaded.json()["id"]
    dataset_id = uploaded.json()["datasets"][0]["id"]

    patched = client.patch(
        f"/api/workspaces/{workspace_id}/sources/{source_id}",
        json={
            "description": "Annual grant awards",
            "metric_hints": {"amount": "Grant amount in INR"},
        },
    )
    assert patched.status_code == 200, patched.text
    assert patched.json()["description"] == "Annual grant awards"
    assert patched.json()["metric_hints"] == {"amount": "Grant amount in INR"}
    profile_before = client.get(f"/api/datasets/{dataset_id}/profile").json()

    with sessions() as session:
        source = session.get(Source, source_id)
        assert source.details["byte_size"] == len(b"id,amount\n1,20.00\n")
        assert source.details["dataset_count"] == 1

    description_only = client.patch(
        f"/api/workspaces/{workspace_id}/sources/{source_id}",
        json={"description": "Updated description"},
    )
    assert description_only.json()["metric_hints"] == {"amount": "Grant amount in INR"}
    assert client.get(f"/api/datasets/{dataset_id}/profile").json() == profile_before

    with sessions() as session:
        other_workspace = Workspace(label="different workspace")
        session.add(other_workspace)
        session.commit()
        other_id = other_workspace.id
    outside_scope = client.patch(
        f"/api/workspaces/{other_id}/sources/{source_id}",
        json={"description": "not allowed"},
    )
    assert outside_scope.status_code == 404

    cleared = client.patch(
        f"/api/workspaces/{workspace_id}/sources/{source_id}",
        json={"description": None, "metric_hints": {}},
    )
    assert cleared.json()["description"] is None
    assert cleared.json()["metric_hints"] == {}


def test_api_rejects_unsupported_and_oversized_files(file_api, monkeypatch):
    client, _, workspace_id, _ = file_api
    response = client.post(
        f"/api/workspaces/{workspace_id}/sources/files",
        files={"file": ("script.py", b"print(1)", "text/plain")},
    )
    assert response.status_code == 415
    monkeypatch.setattr(
        "app.api.files.get_settings",
        lambda: Settings(_env_file=None, max_upload_bytes=4),
    )
    response = client.post(
        f"/api/workspaces/{workspace_id}/sources/files",
        files={"file": ("data.csv", b"a,b\n1,2\n", "text/csv")},
    )
    assert response.status_code == 413
