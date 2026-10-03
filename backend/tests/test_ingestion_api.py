"""Round-trip format uploads and retry/removal behavior without model calls."""

from pathlib import Path
from uuid import UUID
import io

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import select
from pptx import Presentation
from pptx.util import Inches

from app.api.documents import router
from app.api.document_lifecycle import router as lifecycle_router
from app.db.models import Document, Job, Workspace
from app.db.session import get_session
from app.sources.documents import extract_document, process_document
from app.storage.filesystem import FileStorage
from tests.test_documents import _database, _settings


def test_valid_pptx_preserves_distinct_slide_table_ids() -> None:
    deck = Presentation()
    for number in (1, 2):
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = f"Section {number}"
        table = slide.shapes.add_table(
            2, 2, Inches(1), Inches(2), Inches(5), Inches(1)
        ).table
        table.cell(0, 0).text = "Code"
        table.cell(0, 1).text = "Amount"
        table.cell(1, 0).text = f"A{number}"
        table.cell(1, 1).text = str(number * 100)
    stream = io.BytesIO()
    deck.save(stream)
    blocks, _, _ = extract_document("deck.pptx", stream.getvalue())
    tables = [block for block in blocks if block.kind == "table_row"]
    assert {block.location["table"] for block in tables} == {0, 1}
    assert {block.location["slide"] for block in tables} == {1, 2}
    assert tables[-1].location["cells"] == ["A2", "200"]


@pytest.mark.parametrize(
    "name,content,kind",
    [
        ("policy.txt", b"The grant is for eligible families.", "txt"),
        ("policy.md", b"# Policy\nThe grant is for eligible families.", "md"),
        ("policy.html", b"<div>The grant is for eligible families.</div>", "html"),
    ],
)
def test_format_upload_strategy_retry_and_removal(
    name: str,
    content: bytes,
    kind: str,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    sessions, session = _database()
    workspace = Workspace(label="round trip")
    session.add(workspace)
    session.commit()
    workspace_id = workspace.id
    session.close()
    settings = _settings(tmp_path)
    storage = FileStorage(tmp_path)
    monkeypatch.setattr("app.api.documents.get_settings", lambda: settings)
    monkeypatch.setattr("app.api.documents.get_storage", lambda _: storage)
    app = FastAPI()
    app.include_router(router)
    app.include_router(lifecycle_router)

    def db():
        with sessions() as active:
            yield active

    app.dependency_overrides[get_session] = db
    client = TestClient(app)
    path = f"/api/workspaces/{workspace_id}/documents"
    upload = client.post(
        path, files={"file": (name, content)}, data={"chunk_strategy": "parent_child"}
    )
    assert upload.status_code == 202, upload.text
    result = upload.json()
    assert result["source"]["kind"] == kind
    doc_id = result["document"]["id"]
    assert result["document"]["details"]["chunk_strategy"] == "parent_child"
    assert process_document(doc_id, settings, sessions)["state"] == "ready"
    # Identical content plus profile deduplicates, distinct chunk profiles don't.
    duplicate = client.post(
        path, files={"file": (name, content)}, data={"chunk_strategy": "parent_child"}
    )
    assert duplicate.json()["document"]["id"] == doc_id
    other = client.post(
        path, files={"file": (name, content)}, data={"chunk_strategy": "recursive"}
    )
    assert other.json()["document"]["id"] != doc_id
    with sessions() as check:
        document = check.get(Document, doc_id)
        assert document.details["chunk_body_token_limit"] == 240
        document.state = "failed"
        job = check.scalar(
            select(Job).where(Job.payload["document_id"].as_string() == doc_id)
        )
        job.state = "failed"
        check.commit()
    retry = client.post(f"/api/documents/{doc_id}/retry")
    assert retry.status_code == 202
    assert (
        client.post(f"/api/documents/{doc_id}/retry").json()["job_id"]
        == retry.json()["job_id"]
    )
    resumed = process_document(doc_id, settings, sessions)
    assert resumed["extraction_reused"]
    assert client.delete(f"/api/documents/{doc_id}").status_code == 200
    assert client.post(f"/api/documents/{doc_id}/retry").status_code == 404
    assert doc_id not in {document["id"] for document in client.get(path).json()}
    new = client.post(
        path, files={"file": (name, content)}, data={"chunk_strategy": "parent_child"}
    )
    assert new.status_code == 202
    assert new.json()["document"]["id"] != doc_id
    assert (
        client.post(
            path, files={"file": (name, content)}, data={"chunk_strategy": "semantic"}
        ).status_code
        == 422
    )
