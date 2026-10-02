from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest
from docx import Document as WordDocument
from reportlab.pdfgen import canvas  # type: ignore[import-untyped]

from app.config import Settings
from app.ingestion import extractors
from app.sources.documents import (
    DocumentIngestionError,
    ExtractedBlock,
    _make_chunks,
    extract_document,
    validate_document_upload,
)


def _pdf(text: str | None) -> bytes:
    stream = io.BytesIO()
    page = canvas.Canvas(stream)
    if text:
        page.drawString(72, 720, text)
    page.showPage()
    page.save()
    return stream.getvalue()


def test_plain_formats_preserve_headings_and_locations() -> None:
    markdown = b"# Policy\n\n## Eligibility\nApplicants may receive support.\n"
    assert validate_document_upload("guide.md", markdown, 1024) == "md"
    blocks, warnings, _ = extract_document("guide.md", markdown)
    assert not warnings
    assert [(block.kind, block.text) for block in blocks] == [
        ("heading", "Policy"),
        ("heading", "Eligibility"),
        ("paragraph", "Applicants may receive support."),
    ]
    assert blocks[-1].heading == "Policy > Eligibility"
    assert blocks[-1].location["line"] == 4


def test_html_scripts_are_removed_and_table_cells_are_preserved() -> None:
    content = (
        b"<h1>Scheme</h1><p>Eligibility</p><script>private()</script>"
        b"<table><tr><th>Code</th><th>Amount</th></tr>"
        b"<tr><td>A1</td><td>100</td></tr></table>"
    )
    blocks, warnings, _ = extract_document("scheme.html", content)
    assert warnings == ["HTML scripts, styles, and embedded content were omitted."]
    assert all("private" not in block.text for block in blocks)
    row = next(
        block
        for block in blocks
        if block.kind == "table_row" and block.location["row"] == 1
    )
    assert row.location["cells"] == ["A1", "100"]
    assert row.heading == "Scheme"


def test_html_div_body_and_bare_text_are_not_dropped() -> None:
    blocks, warnings, _ = extract_document(
        "content.html",
        b"Root text<div>div text<p>paragraph</p>tail</div><img src='figure.png'>",
    )
    assert [block.text for block in blocks] == [
        "Root text",
        "div text",
        "paragraph",
        "tail",
    ]
    assert warnings == ["HTML embedded images were not OCR processed."]


def test_docx_table_cells_are_available_for_review() -> None:
    document = WordDocument()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Code"
    table.cell(0, 1).text = "Amount"
    table.cell(1, 0).text = "A1"
    table.cell(1, 1).text = "100"
    output = io.BytesIO()
    document.save(output)
    blocks, _, _ = extract_document("table.docx", output.getvalue())
    assert blocks[1].location["cells"] == ["A1", "100"]
    assert blocks[1].location["cell_locations"] == [
        {"row": 1, "column": 0},
        {"row": 1, "column": 1},
    ]


def test_pdf_uses_digital_text_then_routes_empty_page_to_local_ocr(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def fake_ocr(
        *args: object, **kwargs: object
    ) -> tuple[str, None, dict[str, object]]:
        return (
            "यह scanned grant text है",
            None,
            {
                "ocr_confidence": 91.0,
                "ocr_confidence_type": "mean_tesseract_word_confidence_percent",
            },
        )

    monkeypatch.setattr(extractors, "_ocr_page", fake_ocr)
    blocks, warnings, _ = extract_document("scan.pdf", _pdf(None))
    assert "page_1:ocr_table_cells_not_verified" in warnings
    assert len(blocks) == 1
    assert blocks[0].location["extractor"] == "ocr"
    assert blocks[0].location["ocr_confidence"] == 91.0

    blocks, warnings, _ = extract_document(
        "digital.pdf", _pdf("A sufficiently long digital document paragraph.")
    )
    assert not warnings
    assert blocks[0].location["extractor"] == "digital"


def test_ocr_reports_missing_local_assets_explicitly(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(extractors.shutil, "which", lambda _: None)
    text, issue, _ = extractors._ocr_page(b"ignored", 1, "eng+hin", 2)
    assert text == ""
    assert issue == "ocr_unavailable:tesseract_missing"


def test_optional_profiles_and_chunk_strategies_are_explicit() -> None:
    pytest.importorskip("pdfplumber")
    blocks, _, _ = extract_document(
        "test.pdf",
        _pdf("A sufficiently long digital document paragraph."),
        profile="layout",
    )
    assert blocks[0].location["extractor"] == "digital"

    blocks = [
        ExtractedBlock("paragraph", "One section.", "Section", {}, "en-IN", ["Latin"]),
        ExtractedBlock(
            "paragraph", "Second paragraph.", "Section", {}, "en-IN", ["Latin"]
        ),
    ]
    structured = _make_chunks(blocks, strategy="structure")
    recursive = _make_chunks(blocks, strategy="recursive")
    parent_child = _make_chunks(
        [ExtractedBlock("paragraph", "word " * 160, "Section", {}, "en-IN", ["Latin"])],
        strategy="parent_child",
    )
    ordinary = _make_chunks(
        [ExtractedBlock("paragraph", "word " * 160, "Section", {}, "en-IN", ["Latin"])],
        strategy="structure",
    )
    assert len(structured) == 1
    assert len(recursive) == 2
    assert len(parent_child) > len(ordinary)
    assert all(chunk.location["chunk_strategy"] == "recursive" for chunk in recursive)
    with pytest.raises(DocumentIngestionError, match="Semantic chunking"):
        _make_chunks(blocks, strategy="semantic")


def test_pdf_layout_retains_all_wide_table_rows_cells_and_provenance() -> None:
    fixture = Path(__file__).resolve().parents[2] / "evals/fixtures/v2/wide-table.pdf"
    if not fixture.is_file():
        pytest.skip("Phase 04 wide-table fixture has not been generated yet")
    blocks, warnings, _ = extract_document(
        fixture.name, fixture.read_bytes(), ocr_enabled=False, profile="layout"
    )
    rows = [block for block in blocks if block.kind == "table_row"]
    assert not warnings
    assert len(rows) == 15  # Header plus 14 data rows.
    assert all(row.location["cell_count"] == 6 for row in rows)
    assert rows[0].location["cells"] == [
        "application_id",
        "name",
        "district",
        "grant_inr",
        "income_inr",
        "status",
    ]
    assert rows[-1].location["row"] == 14
    assert rows[0].location["review_required"] is True
    assert rows[0].location["cell_locations"][1]["bbox"] == [170.0, 90.0, 300.0, 116.0]


def test_pptx_archive_validation_rejects_unsafe_paths() -> None:
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w") as package:
        package.writestr("../escape", "x")
    with pytest.raises(DocumentIngestionError, match="unsafe path"):
        validate_document_upload("slides.pptx", archive.getvalue(), 1024)


def test_ocr_settings_require_both_local_languages_and_bounded_timeout() -> None:
    with pytest.raises(ValueError, match="both eng and hin"):
        Settings(_env_file=None, ocr_languages="eng")
    with pytest.raises(ValueError, match="cannot exceed 60"):
        Settings(_env_file=None, ocr_timeout_seconds=61)


def test_ocr_missing_hindi_assets_does_not_render_or_fallback(monkeypatch):
    monkeypatch.setattr(extractors.shutil, "which", lambda _: "/fake/tesseract")
    monkeypatch.setattr(extractors, "_tesseract_languages", lambda _: {"eng", "osd"})
    text, issue, metadata = extractors._ocr_page(b"ignored", 1, "eng+hin", 2)
    assert text == "" and metadata == {}
    assert issue == "ocr_unavailable:missing_language_assets:hin"


def test_ocr_timeout_returns_explicit_page_issue(monkeypatch):
    import subprocess

    monkeypatch.setattr(extractors.shutil, "which", lambda _: "/fake/tesseract")
    monkeypatch.setattr(
        extractors, "_tesseract_languages", lambda _: {"eng", "hin", "osd"}
    )
    monkeypatch.setattr(extractors, "_pdfium_page_png", lambda *args: b"png")

    def timeout(*args, **kwargs):
        raise subprocess.TimeoutExpired("tesseract", 2)

    monkeypatch.setattr(extractors.subprocess, "run", timeout)
    text, issue, metadata = extractors._ocr_page(b"ignored", 1, "eng+hin", 2)
    assert text == "" and issue == "ocr_timeout" and metadata == {}


def test_oversized_pdf_page_rejected_before_allocating_and_resources_close(monkeypatch):
    import pypdfium2 as pdfium

    closed = []

    class Page:
        def get_size(self):
            return (10000000, 10000000)

        def render(self, **kwargs):
            pytest.fail("Oversized page must not allocate a bitmap")

        def close(self):
            closed.append("page")

    class Document:
        def __getitem__(self, index):
            return Page()

        def close(self):
            closed.append("document")

    monkeypatch.setattr(pdfium, "PdfDocument", lambda _: Document())
    with pytest.raises(DocumentIngestionError) as error:
        extractors._pdfium_page_png(b"ignored", 1)
    assert error.value.code == "pdf_render_dimensions_exceeded"
    assert closed == ["page", "document"]
