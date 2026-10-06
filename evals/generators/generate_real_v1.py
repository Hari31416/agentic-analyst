"""Build a reproducible, source-backed real-data evaluation inventory.

Run from the repository root with ``cd backend && uv run python
../evals/generators/generate_real_v1.py`` after placing the source files under
``evals/fixtures/real-v1/raw``. Labels are calculated from the source workbook
with Decimal arithmetic; survey passage labels are checked against extracted
PDF text before they are emitted.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from collections import defaultdict
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[2]
FIXTURE_ROOT = ROOT / "evals" / "fixtures"
PACK_ROOT = FIXTURE_ROOT / "real-v1"
RAW_ROOT = PACK_ROOT / "raw"
GENERATED_ROOT = PACK_ROOT / "derived"
WORKBOOK = RAW_ROOT / "online-retail.xlsx"
RAW_SURVEYS = {
    "survey-en": RAW_ROOT / "economic-survey-2023-24-en.pdf",
    "survey-hi": RAW_ROOT / "economic-survey-2023-24-hi.pdf",
    "inflation-en": RAW_ROOT / "inflation-en.pdf",
    "inflation-hi": RAW_ROOT / "inflation-hi.pdf",
    "appendix-en": RAW_ROOT / "statistical-appendix-en.pdf",
    "appendix-hi": RAW_ROOT / "statistical-appendix-hi.pdf",
}
MONTH_CSV = GENERATED_ROOT / "online-retail-2011-01.csv"
RULES = GENERATED_ROOT / "retail-analysis-rules.txt"
VERSION = "real-v1"
ALLOWED_ACTIONS = [
    "finish_answer",
    "list_sources",
    "dataset_profile",
    "inspect_schema",
    "sample_rows",
    "run_sql",
    "run_python",
    "analyze_data",
    "inspect_artifact",
    "list_artifacts",
    "search_documents",
    "source_passage",
    "summarize_documents",
]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def require_file(path: Path, *, required: bool = True) -> bool:
    if path.is_file():
        return True
    if required:
        raise FileNotFoundError(f"Required source file is missing: {path}")
    return False


def verify_download_manifest() -> dict[str, dict[str, Any]]:
    manifest_path = PACK_ROOT / "downloads.json"
    require_file(manifest_path)
    document = json.loads(manifest_path.read_text(encoding="utf-8"))
    entries = {entry["name"]: entry for entry in document.get("files", [])}
    required = {"online-retail.xlsx", *(path.name for path in RAW_SURVEYS.values())}
    if not required.issubset(entries):
        raise ValueError(
            f"Download manifest is missing: {sorted(required - set(entries))}"
        )
    for name in sorted(required):
        path = RAW_ROOT / name
        require_file(path)
        entry = entries[name]
        if path.stat().st_size != entry["bytes"] or sha256(path) != entry["sha256"]:
            raise ValueError(
                f"Downloaded source failed pinned size/hash verification: {name}"
            )
    return entries


def decimal_value(value: Any) -> Decimal:
    if value is None or isinstance(value, bool):
        return Decimal(0)
    return Decimal(str(value))


def normalize_header(value: Any) -> str:
    return str(value).strip()


def write_retail_slice() -> tuple[list[dict[str, str]], dict[str, Any]]:
    require_file(WORKBOOK)
    GENERATED_ROOT.mkdir(parents=True, exist_ok=True)
    book = load_workbook(WORKBOOK, read_only=True, data_only=True)
    sheet = book.active
    rows = sheet.iter_rows(values_only=True)
    headers = [normalize_header(value) for value in next(rows)]
    required = {
        "InvoiceNo",
        "StockCode",
        "Description",
        "Quantity",
        "InvoiceDate",
        "UnitPrice",
        "CustomerID",
        "Country",
    }
    if not required.issubset(headers):
        book.close()
        raise ValueError(
            f"Online Retail columns missing: {sorted(required - set(headers))}"
        )
    selected: list[dict[str, str]] = []
    for values in rows:
        record = dict(zip(headers, values))
        invoice_date = record.get("InvoiceDate")
        if (
            isinstance(invoice_date, datetime)
            and invoice_date.year == 2011
            and invoice_date.month == 1
        ):
            selected.append(
                {header: _csv_value(record.get(header)) for header in headers}
            )
    book.close()
    if not selected:
        raise ValueError("No January 2011 records found in the workbook")
    with MONTH_CSV.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=headers, lineterminator="\n")
        writer.writeheader()
        writer.writerows(selected)
    return selected, {"row_count": len(selected), "columns": headers}


def _csv_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, float):
        return str(Decimal(str(value)))
    return str(value)


def qualifying_rows(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    return [
        row
        for row in rows
        if not row["InvoiceNo"].strip().upper().startswith("C")
        and decimal_value(row["Quantity"]) > 0
        and decimal_value(row["UnitPrice"]) > 0
    ]


def retail_gold(rows: list[dict[str, str]]) -> dict[str, Any]:
    kept = qualifying_rows(rows)
    sales_by_country: dict[str, Decimal] = defaultdict(Decimal)
    customer_ids = set()
    missing_customer_lines = 0
    for row in kept:
        sales = decimal_value(row["Quantity"]) * decimal_value(row["UnitPrice"])
        sales_by_country[row["Country"]] += sales
        customer = row["CustomerID"].strip()
        if customer:
            customer_ids.add(customer)
        else:
            missing_customer_lines += 1
    total_sales = sum(sales_by_country.values(), Decimal(0))
    units = sum((decimal_value(row["Quantity"]) for row in kept), Decimal(0))
    invoice_count = len({row["InvoiceNo"] for row in kept})
    country_count = len(sales_by_country)
    top_country, top_sales = min(
        sales_by_country.items(), key=lambda item: (-item[1], item[0])
    )
    return {
        "qualifying_rows": len(kept),
        "sales_gbp": total_sales,
        "qualifying_units": units,
        "invoice_count": invoice_count,
        "country_count": country_count,
        "customer_count": len(customer_ids),
        "missing_customer_lines": missing_customer_lines,
        "missing_customer_percent": Decimal(missing_customer_lines)
        * 100
        / Decimal(len(kept)),
        "top_country": top_country,
        "top_country_sales_gbp": top_sales,
        "top_five_countries": sorted(
            sales_by_country.items(), key=lambda item: (-item[1], item[0])
        )[:5],
        "missing_customer_sales_gbp": sum(
            (
                decimal_value(row["Quantity"]) * decimal_value(row["UnitPrice"])
                for row in kept
                if not row["CustomerID"].strip()
            ),
            Decimal(0),
        ),
        "country_sales": dict(sales_by_country),
        "cancelled_rows": sum(
            row["InvoiceNo"].strip().upper().startswith("C") for row in rows
        ),
        "negative_quantity_rows": sum(
            decimal_value(row["Quantity"]) < 0 for row in rows
        ),
        "negative_price_rows": sum(decimal_value(row["UnitPrice"]) < 0 for row in rows),
    }


def write_rules() -> None:
    RULES.parent.mkdir(parents=True, exist_ok=True)
    RULES.write_text(
        "Retail evaluation rules (real-v1)\n\n"
        "The attached CSV is a row-preserving slice of the UCI Online Retail "
        "workbook containing every source row whose InvoiceDate falls in "
        "January 2011. InvoiceNo is an invoice identifier; rows are invoice "
        "lines and should not be treated as unique invoices.\n\n"
        "For qualifying sales questions, include a row only when InvoiceNo does not "
        "begin with C (case-insensitive), Quantity is greater than zero, and "
        "UnitPrice is greater than zero. Sales are Quantity multiplied by "
        "UnitPrice and are denominated in GBP. Preserve missing CustomerID as "
        "missing; do not infer identities. Dates in this slice use the "
        "workbook's local timestamps. These rules define an evaluation "
        "convention and do not claim to be UCI's official accounting policy.\n",
        encoding="utf-8",
    )


def source(
    alias: str, name: str, path: Path, kind: str, version: str = "1"
) -> dict[str, Any]:
    # Each runner trial uploads a new immutable source, whose API version is 1.
    # The pack release identifier belongs in the inventory/manifest, not here.
    return {
        "alias": alias,
        "name": name,
        "kind": kind,
        "path": path.relative_to(FIXTURE_ROOT).as_posix(),
        "sha256": sha256(path),
        "version": version,
    }


def case(
    ident: str,
    question: str,
    language: str,
    tags: list[str],
    answerability: str,
    sources: list[dict[str, Any]],
    *,
    calculations: list[dict[str, Any]] | None = None,
    passages: list[dict[str, str]] | None = None,
    artifacts: list[dict[str, Any]] | None = None,
    rubric: dict[str, Any] | None = None,
    answer_language: str | None = None,
) -> dict[str, Any]:
    # The current scorer reads the last successful structured result and uses
    # column names to identify scalars. Make this execution contract explicit
    # rather than failing a correct answer that chose different SQL aliases.
    if calculations:
        columns = ", ".join(item.get("column") or item["key"] for item in calculations)
        if language == "hi-IN":
            question += (
                " अंतिम संख्यात्मक परिणाम एक ही SQL/analysis परिणाम पंक्ति में दें। "
                f"इन कॉलम नामों का उपयोग करें: {columns}।"
            )
        elif language.startswith("hi-Latn"):
            question += (
                " Final numeric results ek hi SQL/analysis result row mein dein, "
                f"in column names ke saath: {columns}."
            )
        else:
            question += (
                " Return the final numeric results together in one SQL/analysis "
                f"result row with these column names: {columns}."
            )
    result = {
        "schema_version": 1,
        "id": ident,
        "question": question,
        "language": language,
        "tags": tags,
        "answerability": answerability,
        "sources": sources,
        "expectations": {
            "allowed_actions": ALLOWED_ACTIONS,
            "calculations": calculations or [],
            "passages": passages or [],
            "artifacts": artifacts or [],
            "rubric": rubric or {},
        },
        "review": {"provenance": "unreviewed"},
    }
    if answer_language is not None:
        result["answer_language"] = answer_language
    return result


def calc(
    key: str,
    value: Any,
    unit: str | None = None,
    tolerance: str = "0",
    column: str | None = None,
) -> dict[str, Any]:
    result: dict[str, Any] = {
        "key": key,
        "value": str(value),
        "unit": unit,
        "tolerance": tolerance,
    }
    if column:
        result["column"] = column
    return result


def extract_pdf_text(path: Path) -> str:
    reader = PdfReader(path)
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def normalized(value: str) -> str:
    return re.sub(r"\s+", " ", value).casefold().strip()


def survey_case(
    *,
    ident: str,
    question: str,
    language: str,
    tags: list[str],
    path: Path,
    alias: str,
    expected_phrase: str,
    note: str,
) -> dict[str, Any]:
    text = extract_pdf_text(path)
    if normalized(expected_phrase) not in normalized(text):
        raise ValueError(
            f"Survey passage label was not found in {path.name}: {expected_phrase!r}"
        )
    return case(
        ident,
        question,
        language,
        tags,
        "answerable",
        [source(alias, path.name, path, "pdf")],
        passages=[{"source_alias": alias, "contains": expected_phrase}],
        rubric={
            "manual_review": note,
            "answer_requirements": "Check factual scope, year, original units, page citation, and language fidelity against the cited passage.",
        },
    )


def create_cases(
    rows: list[dict[str, str]],
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    gold = retail_gold(rows)
    retail_src = source("retail", MONTH_CSV.name, MONTH_CSV, "csv")
    rules_src = source("rules", RULES.name, RULES, "txt")
    cases: list[dict[str, Any]] = []

    def add(*args: Any, **kwargs: Any) -> None:
        cases.append(case(*args, **kwargs))

    add(
        "retail-jan-sales-en",
        "Using the attached retail-analysis rules, return one aggregate row with columns sales_gbp, qualifying_rows, and invoice_count for January 2011. Cite the inclusion rules and calculation evidence and state GBP.",
        "en-IN",
        ["retail", "sql", "aggregation"],
        "answerable",
        [retail_src, rules_src],
        calculations=[
            calc("sales_gbp", gold["sales_gbp"], "GBP", "0.01"),
            calc("qualifying_rows", gold["qualifying_rows"]),
            calc("invoice_count", gold["invoice_count"]),
        ],
        passages=[
            {"source_alias": "rules", "contains": "Quantity is greater than zero"}
        ],
        rubric={
            "manual_review": "Check the stated row filters and invoice-line versus invoice grain. 'Qualifying sales' is the evaluation convention, not source accounting policy."
        },
    )
    add(
        "retail-distinct-countries",
        "Apply the attached rules and return country_count, the number of distinct countries with qualifying sales in January 2011.",
        "en-IN",
        ["retail", "distinct", "sql"],
        "answerable",
        [retail_src, rules_src],
        calculations=[calc("country_count", gold["country_count"])],
        rubric={"manual_review": "Distinct country count after row filters."},
    )
    add(
        "retail-customer-null-en",
        "Among qualifying January 2011 invoice lines, return missing_customer_lines and missing_customer_percent. Use all qualifying lines as the percentage denominator and round the percentage to two decimal places.",
        "en-IN",
        ["retail", "nulls", "aggregation"],
        "answerable",
        [retail_src, rules_src],
        calculations=[
            calc("missing_customer_lines", gold["missing_customer_lines"]),
            calc(
                "missing_customer_percent",
                gold["missing_customer_percent"].quantize(Decimal("0.01")),
                tolerance="0.01",
            ),
        ],
        rubric={
            "manual_review": "Denominator is qualifying invoice lines, not distinct invoices or all raw rows. State percentage units."
        },
    )
    add(
        "retail-top-country-en",
        "Which country has the highest qualifying sales in January 2011? Return a one-row aggregate with columns top_country and top_country_sales_gbp.",
        "en-IN",
        ["retail", "groupby", "ranking"],
        "answerable",
        [retail_src, rules_src],
        calculations=[
            calc("top_country_sales_gbp", gold["top_country_sales_gbp"], "GBP", "0.01")
        ],
        rubric={
            "expected_category": gold["top_country"],
            "manual_review": "The runner checks the sales amount; manually check the country category and its ordering.",
        },
    )
    add(
        "retail-cancellations",
        "Return cancelled_rows for all raw January 2011 rows whose invoice number begins with C. In the same response return qualifying sales as sales_gbp after applying all attached rules.",
        "en-IN",
        ["retail", "filters", "cancellations"],
        "answerable",
        [retail_src, rules_src],
        calculations=[
            calc("cancelled_rows", gold["cancelled_rows"]),
            calc("sales_gbp", gold["sales_gbp"], "GBP", "0.01"),
        ],
        rubric={
            "manual_review": "Cancellation count uses raw month rows; sales applies all three rules."
        },
    )
    add(
        "retail-sales-chart-csv",
        "Calculate qualifying sales by country for January 2011. Save all country groups as CSV and create a labelled bar chart of the top five countries in GBP.",
        "en-IN",
        ["retail", "artifact", "chart", "csv"],
        "answerable",
        [retail_src, rules_src],
        artifacts=[{"media_type": "text/csv"}, {"media_type": "image/png"}],
        rubric={
            "expected_ordered_country_sales": [
                [country, str(value)] for country, value in gold["top_five_countries"]
            ],
            "manual_review": "Artifact type/presence is checked by the runner; inspect CSV values, top-five order, and chart labels manually.",
        },
    )
    add(
        "retail-revenue-hindi",
        "संलग्न नियमों के अनुसार जनवरी 2011 में qualifying sales निकालें। एक aggregate row में sales_gbp, qualifying_rows और invoice_count दें।",
        "hi-IN",
        ["retail", "hindi", "aggregation"],
        "answerable",
        [retail_src, rules_src],
        calculations=[
            calc("sales_gbp", gold["sales_gbp"], "GBP", "0.01"),
            calc("qualifying_rows", gold["qualifying_rows"]),
            calc("invoice_count", gold["invoice_count"]),
        ],
        rubric={
            "manual_review": "Native-speaker review is required for Hindi correctness and terminology."
        },
    )
    add(
        "retail-codeswitch-top-country",
        "Jan 2011 mein rules apply karke highest qualifying sales wala country batayein; us country ki sales GBP mein aur sabhi qualifying rows se overall distinct invoice count bhi dein.",
        "hi-Latn-IN",
        ["retail", "code-switch", "aggregation"],
        "answerable",
        [retail_src, rules_src],
        calculations=[
            calc("top_country_sales_gbp", gold["top_country_sales_gbp"], "GBP", "0.01"),
            calc("invoice_count", gold["invoice_count"]),
        ],
        rubric={
            "expected_category": gold["top_country"],
            "manual_review": "Transliterated Hindi/English code-switching needs human review.",
        },
        answer_language="hi-IN",
    )
    add(
        "retail-mixed-rule-retrieval",
        "Find the row inclusion rules in the attached note, then return a one-row aggregate with sales_gbp for January 2011. Cite the rules and calculation evidence.",
        "en-IN",
        ["retail", "mixed-source", "retrieval", "calculation"],
        "answerable",
        [rules_src, retail_src],
        calculations=[calc("sales_gbp", gold["sales_gbp"], "GBP", "0.01")],
        passages=[
            {"source_alias": "rules", "contains": "InvoiceNo does not begin with C"}
        ],
        rubric={"manual_review": "Verify the retrieved rule is applied to the data."},
    )
    add(
        "retail-missing-cost-unsupported",
        "What were gross profit and profit margin in January 2011?",
        "en-IN",
        ["retail", "unsupported", "abstention"],
        "unsupported",
        [retail_src],
        rubric={
            "manual_review": "The runner checks clarification status; manually verify no invented profit/margin is given."
        },
    )
    add(
        "retail-lines-vs-invoices",
        "Return distinct invoice_count and qualifying_rows for January 2011 under the attached rules. Briefly explain the difference between invoice count and invoice-line count.",
        "en-IN",
        ["retail", "grain", "distinct"],
        "answerable",
        [retail_src, rules_src],
        calculations=[
            calc("invoice_count", gold["invoice_count"]),
            calc("qualifying_rows", gold["qualifying_rows"]),
        ],
        rubric={
            "manual_review": "Check distinction between row grain and invoice grain."
        },
    )
    add(
        "retail-units-sold",
        "Using the attached rules, return the sum as qualifying_units: total quantity from qualifying January 2011 invoice lines.",
        "en-IN",
        ["retail", "sum", "units"],
        "answerable",
        [retail_src, rules_src],
        calculations=[calc("qualifying_units", gold["qualifying_units"])],
        rubric={
            "manual_review": "Check filters and describe the result as items/units."
        },
    )
    add(
        "retail-missing-customer-sales",
        "Return missing_customer_sales_gbp and sales_gbp for January 2011 under the attached rules. Also return missing_customer_sales_percent, the percentage of qualifying sales value from rows with a missing CustomerID, rounded to two decimal places.",
        "en-IN",
        ["retail", "nulls", "ratio"],
        "answerable",
        [retail_src, rules_src],
        calculations=[
            calc(
                "missing_customer_sales_gbp",
                gold["missing_customer_sales_gbp"],
                "GBP",
                "0.01",
            ),
            calc("sales_gbp", gold["sales_gbp"], "GBP", "0.01"),
            calc(
                "missing_customer_sales_percent",
                (gold["missing_customer_sales_gbp"] * 100 / gold["sales_gbp"]).quantize(
                    Decimal("0.01")
                ),
                tolerance="0.01",
            ),
        ],
        rubric={
            "manual_review": "Do not impute customer identity. Check the ratio uses qualifying sales value."
        },
    )
    add(
        "retail-top-five",
        "Return the five countries with the largest qualifying sales in January 2011, ordered highest to lowest, with each value in GBP.",
        "en-IN",
        ["retail", "ranking", "top-k"],
        "answerable",
        [retail_src, rules_src],
        rubric={
            "expected_ordered_country_sales": [
                [country, str(value)] for country, value in gold["top_five_countries"]
            ],
            "manual_review": "The case schema cannot score per-category grouped values; compare names, order, and amounts with manifest.",
        },
    )
    add(
        "retail-hindi-english-rules",
        "Apply the attached English rules to January 2011. योग्य rows की count और sales_gbp बताइए; missing CustomerID को missing ही रखें.",
        "hi-IN",
        ["retail", "code-switch", "filters", "nulls"],
        "answerable",
        [retail_src, rules_src],
        calculations=[
            calc("qualifying_rows", gold["qualifying_rows"]),
            calc("sales_gbp", gold["sales_gbp"], "GBP", "0.01"),
        ],
        rubric={
            "manual_review": "Review transliterated/code-switched language and missing-value treatment."
        },
    )
    add(
        "retail-export-and-chart-hi",
        "संलग्न नियमों के अनुसार जनवरी 2011 में देश के अनुसार qualifying sales निकालें। सभी देशों के परिणाम CSV में सहेजें और शीर्ष पाँच देशों का GBP में नामांकित bar chart बनाएँ।",
        "hi-IN",
        ["retail", "hindi", "artifact", "chart", "csv"],
        "answerable",
        [retail_src, rules_src],
        artifacts=[{"media_type": "text/csv"}, {"media_type": "image/png"}],
        rubric={
            "manual_review": "Inspect exported values and chart labels; language quality needs a native speaker."
        },
    )
    add(
        "retail-cancellation-only-count",
        "Return cancelled_rows: the count of raw January 2011 rows whose InvoiceNo begins with C, without applying other filters.",
        "en-IN",
        ["retail", "filter", "cancellations"],
        "answerable",
        [retail_src],
        calculations=[calc("cancelled_rows", gold["cancelled_rows"])],
        rubric={"manual_review": "Count raw rows; no other filters apply."},
    )
    add(
        "retail-distinct-customer-count",
        "After applying the attached rules, return customer_count, the count of distinct non-missing customers with qualifying January 2011 invoice lines.",
        "en-IN",
        ["retail", "distinct", "nulls"],
        "answerable",
        [retail_src, rules_src],
        calculations=[calc("customer_count", gold["customer_count"])],
        rubric={"manual_review": "Distinct count excludes missing CustomerID."},
    )
    add(
        "retail-netherlands-sales",
        "Using the attached rules, calculate qualifying January 2011 sales for Country = Netherlands. Return one aggregate value in a column named sales_gbp.",
        "en-IN",
        ["retail", "filtered-aggregation", "country"],
        "answerable",
        [retail_src, rules_src],
        calculations=[
            calc(
                "sales_gbp",
                gold["country_sales"].get("Netherlands", Decimal(0)),
                "GBP",
                "0.01",
            )
        ],
        rubric={
            "manual_review": "Verify country filter and the three inclusion rules."
        },
    )
    add(
        "retail-scope-clarification",
        "Give total qualifying sales for all of 2011.",
        "en-IN",
        ["retail", "ambiguous", "scope"],
        "ambiguous",
        [retail_src, rules_src],
        rubric={
            "manual_review": "Verify temporal coverage is clarified; scorer only checks clarification status."
        },
    )

    inflation_en = RAW_SURVEYS["inflation-en"]
    inflation_hi = RAW_SURVEYS["inflation-hi"]
    survey_en = RAW_SURVEYS["survey-en"]
    appendix_en = RAW_SURVEYS["appendix-en"]
    en_chapter = source("inflation", inflation_en.name, inflation_en, "pdf")
    hi_chapter = source("inflation-hi", inflation_hi.name, inflation_hi, "pdf")
    en_survey = source("survey", survey_en.name, survey_en, "pdf")
    appendix = source("appendix", appendix_en.name, appendix_en, "pdf")
    cases.append(
        survey_case(
            ident="survey-fy24-retail-inflation-en",
            question="What retail inflation did India record in FY24, and how does the Survey characterize that level relative to the Covid-19 pandemic? Cite the chapter passage and preserve the percentage unit.",
            language="en-IN",
            tags=["survey", "inflation", "numeric-claim", "retrieval"],
            path=inflation_en,
            alias="inflation",
            expected_phrase="retail inflation at 5.4 per cent in FY24",
            note="The 5.4 per cent claim is supported by this phrase; manually verify page/paragraph and context.",
        )
    )
    add(
        "survey-fy24-inflation-hi-query-en-source",
        "FY24 में भारत की खुदरा मुद्रास्फीति कितनी थी? Economic Survey के अंग्रेज़ी अध्याय से साक्ष्य दें और प्रतिशत इकाई बनाए रखें।",
        "hi-IN",
        ["survey", "inflation", "hindi-query", "cross-language"],
        "answerable",
        [en_chapter],
        passages=[
            {
                "source_alias": "inflation",
                "contains": "retail inflation at 5.4 per cent in FY24",
            }
        ],
        rubric={
            "manual_review": "Verify Hindi meaning, 5.4 per cent, FY24 scope, and citation against English source."
        },
    )
    cases.append(
        survey_case(
            ident="survey-inflation-drivers-en",
            question="What factors does the opening summary associate with core price pressures in FY22/FY23 and food-price pressures in FY23/FY24? Cite evidence and distinguish the two time periods.",
            language="en-IN",
            tags=["survey", "inflation", "causal-language", "retrieval"],
            path=inflation_en,
            alias="inflation",
            expected_phrase="Food prices were affected by adverse weather conditions in the last two years",
            note="Review all opening-page factors without overclaiming causation.",
        )
    )
    cases.append(
        survey_case(
            ident="survey-inflation-policy-measures-en",
            question="Which government measures does the opening summary of the chapter say helped mitigate food inflation? Cite that passage and list all the measures named there.",
            language="en-IN",
            tags=["survey", "inflation", "list-extraction"],
            path=inflation_en,
            alias="inflation",
            expected_phrase="dynamic stock management, open market operations",
            note="Verify the entire list against the full sentence.",
        )
    )
    add(
        "survey-inflation-hi-source",
        "इस हिंदी अध्याय के अनुसार FY24 में खुदरा मुद्रास्फीति और CFPI पर आधारित खाद्य मुद्रास्फीति कितनी थी? FY23 से FY24 तक खाद्य मुद्रास्फीति में बदलाव भी बताएं और मूल साक्ष्य दें।",
        "hi-IN",
        ["survey", "inflation", "hindi-source", "ocr-descriptive"],
        "answerable",
        [hi_chapter],
        rubric={
            "manual_review": "Hindi PDF is visually legible, but embedded extraction uses legacy glyphs. Review rendered page and cited evidence manually; no garbled-text gold label."
        },
    )
    add(
        "survey-summary-vs-chapter",
        "Compare the inflation chapter in the full Economic Survey PDF with the standalone inflation chapter on the FY24 retail inflation figure. Cite both sources and explain whether they agree.",
        "en-IN",
        ["survey", "mixed-source", "cross-document", "retrieval"],
        "answerable",
        [en_survey, en_chapter],
        passages=[
            {
                "source_alias": "survey",
                "contains": "retail inflation at 5.4 per cent in FY24",
            },
            {
                "source_alias": "inflation",
                "contains": "retail inflation at 5.4 per cent in FY24",
            },
        ],
        rubric={
            "manual_review": "Verify both documents are cited for the same FY24 measure and period."
        },
    )
    add(
        "survey-inflation-hi-cross-language",
        "English inflation chapter में FY24 के core inflation के बारे में क्या निष्कर्ष है? उद्धरण और chapter page दें।",
        "hi-IN",
        ["survey", "hindi-query", "code-switch", "cross-language"],
        "answerable",
        [en_chapter],
        passages=[
            {
                "source_alias": "inflation",
                "contains": "core inflation to a four-year low in FY24",
            }
        ],
        rubric={
            "manual_review": "Human review required for code-switched Hindi and exact core-inflation claim."
        },
    )
    add(
        "survey-hindi-source-search-en",
        "Search the Hindi prices and inflation chapter of the Economic Survey, then report one clearly supported FY24 statement in English with a chapter PDF page citation.",
        "en-IN",
        ["survey", "hindi-source", "cross-language", "ocr-descriptive"],
        "answerable",
        [hi_chapter],
        rubric={
            "manual_review": "The 18-page Hindi chapter requires OCR for legacy fonts. Verify the FY24 claim, English translation and chapter PDF page citation against the rendered original."
        },
    )
    cases.append(
        survey_case(
            ident="survey-appendix-table-retrieval",
            question="In the statistical appendix, locate Table 4.3, All India Consumer Price Index Numbers. Identify the table's index families and their coverage categories and report both PDF and printed source page numbers; do not calculate a trend unless the correct column and period are clear.",
            language="en-IN",
            tags=["survey", "appendix", "table-retrieval", "ambiguity"],
            path=appendix_en,
            alias="appendix",
            expected_phrase="Table 4.3. All India Consumer Price Index Numbers",
            note="PDF page 92 starts printed Statistical Appendix page 85; inspect continuation pages and column headers visually.",
        )
    )
    add(
        "survey-causal-overclaim",
        "Does the chapter establish that adverse weather alone caused India's FY24 food inflation? Cite evidence for your conclusion.",
        "en-IN",
        ["survey", "inflation", "causality", "guardrail"],
        "answerable",
        [en_chapter],
        passages=[
            {
                "source_alias": "inflation",
                "contains": "Food prices were affected by adverse weather conditions",
            }
        ],
        rubric={
            "manual_review": "Check the answer says 'alone' is not established by this sentence and separates attribution from causal proof."
        },
    )
    complete_references(cases, gold)
    return cases, gold


def complete_references(cases: list[dict[str, Any]], gold: dict[str, Any]) -> None:
    """Bind complete manual answer keys and diagnostic anchors to each case."""
    document_references = {
        "survey-fy24-retail-inflation-en": {
            "required_claims": [
                "Retail inflation was 5.4% in FY24",
                "Lowest level since the Covid-19 pandemic period",
            ],
            "locations": {"inflation": {"pdf_pages": [1], "paragraphs": ["3.1"]}},
        },
        "survey-fy24-inflation-hi-query-en-source": {
            "required_claims": ["Retail inflation was 5.4% in FY24"],
            "locations": {"inflation": {"pdf_pages": [1], "paragraphs": ["3.1"]}},
        },
        "survey-inflation-drivers-en": {
            "required_claims": [
                "Pandemic supply disruptions and conflict-related commodity prices contributed to core price pressures in FY22/FY23",
                "Adverse weather affected food prices in FY23/FY24",
            ],
            "locations": {
                "inflation": {
                    "pdf_pages": [1, 9],
                    "paragraphs": ["opening summary", "3.18"],
                }
            },
        },
        "survey-inflation-policy-measures-en": {
            "required_claims": [
                "dynamic stock management",
                "open market operations",
                "subsidised provision of essential food items",
                "trade policy measures",
            ],
            "locations": {
                "inflation": {"pdf_pages": [1], "paragraphs": ["opening summary"]}
            },
        },
        "survey-inflation-hi-source": {
            "required_claims": [
                "FY24 retail inflation was 5.4%",
                "FY24 CFPI food inflation was 7.5%, up from 6.6% in FY23",
            ],
            "locations": {
                "inflation-hi": {"pdf_pages": [1, 9], "paragraphs": ["3.1", "3.18"]}
            },
            "measure_distinction": "CFPI food inflation is distinct from the broader food-and-beverages CPI group.",
        },
        "survey-summary-vs-chapter": {
            "required_claims": [
                "Both sources report retail inflation of 5.4% in FY24",
                "Both sources are cited for that same measure and period",
            ],
            "locations": {
                "survey": {"pdf_pages": [132], "paragraphs": ["3.1"]},
                "inflation": {"pdf_pages": [1], "paragraphs": ["3.1"]},
            },
        },
        "survey-inflation-hi-cross-language": {
            "required_claims": ["Core inflation reached a four-year low in FY24"],
            "locations": {
                "inflation": {
                    "pdf_pages": [1, 5],
                    "paragraphs": ["opening summary", "3.11 onward"],
                }
            },
            "period_distinction": "June 2024's 3.1% is not the annual FY24 figure.",
        },
        "survey-hindi-source-search-en": {
            "required_claims": [
                "At least one fact supported by the selected Hindi source",
                "The fact explicitly concerns FY24",
                "Faithful English translation and correct chapter PDF page citation",
            ],
            "accepted_examples": [
                {
                    "claim": "Retail inflation was 5.4% in FY24",
                    "pdf_page": 1,
                    "paragraph": "3.1",
                },
                {
                    "claim": "CFPI food inflation rose from 6.6% in FY23 to 7.5% in FY24",
                    "pdf_page": 9,
                    "paragraph": "3.18",
                },
            ],
            "alternatives": "Any other visually verified FY24 claim is acceptable; examples are not an exhaustive answer key.",
        },
        "survey-appendix-table-retrieval": {
            "required_claims": [
                "CPI-IW: General",
                "CPI-NS: Rural, Urban, Combined",
                "CPI-AL: General",
                "CPI-RL: General",
            ],
            "locations": {
                "appendix": {"pdf_pages": [92, 93], "printed_pages": [85, 86]}
            },
            "pitfalls": [
                "Do not place Rural/Urban/Combined under CPI-IW",
                "Do not infer comparable trends across base-year changes",
            ],
        },
        "survey-causal-overclaim": {
            "required_claims": [
                "The chapter does not establish adverse weather as the sole proven cause",
                "Weather is a reported contributing condition, distinct from proof of sole causation",
            ],
            "locations": {
                "inflation": {
                    "pdf_pages": [1, 9],
                    "paragraphs": ["opening summary", "3.18"],
                }
            },
            "alternatives": "A supported explanation from either passage suffices; listing every condition is not required.",
        },
    }
    for item in cases:
        rubric = item["expectations"]["rubric"]
        rubric["review_criteria"] = {
            "complete": "All requested outputs are correct, supported, in scope, and satisfy the required language/artifact contract.",
            "partial": "Some requested outputs are correct but at least one is missing or incomplete; explain each omission.",
            "failed": "The central result is incorrect, unsupported or absent; record contradictions explicitly.",
            "separate_checks": "Record factual task success, tool appropriateness, execution, citation validity and language quality separately. Preserve deterministic failures.",
        }
        if item["id"] in document_references:
            rubric["reference_answer"] = document_references[item["id"]]
            for passage in item["expectations"]["passages"]:
                passage["match_policy"] = "diagnostic"
        if item["id"] in {"retail-sales-chart-csv", "retail-export-and-chart-hi"}:
            rubric["expected_country_sales"] = {
                country: str(value)
                for country, value in sorted(gold["country_sales"].items())
            }
            rubric["expected_csv_rows"] = len(gold["country_sales"])
            rubric["expected_ordered_country_sales"] = [
                [country, str(value)] for country, value in gold["top_five_countries"]
            ]
            rubric["artifact_requirements"] = (
                "CSV contains all 22 country groups exactly once with GBP values within 0.01; bar chart shows the top five in descending order, with country labels and a GBP axis. Extra execution artifacts do not replace requested outputs."
            )
        if item["id"] == "retail-missing-cost-unsupported":
            rubric["reference_answer"] = {
                "required_claims": [
                    "Sales data does not include costs needed for gross profit or margin",
                    "Request cost data; do not invent costs or margins",
                ]
            }
        if item["id"] == "retail-scope-clarification":
            rubric["reference_answer"] = {
                "required_claims": [
                    "Selected data covers January only",
                    "Ask for remaining months or confirmation to report January; do not extrapolate an annual total",
                ]
            }
        if item["id"] == "retail-lines-vs-invoices":
            rubric["reference_answer"] = {
                "required_claims": [
                    "Invoices are distinct InvoiceNo values; lines are qualifying rows, with multiple lines possible per invoice"
                ]
            }
        if item["id"] == "survey-causal-overclaim":
            item["expectations"]["passages"][0]["contains_any"] = [
                "lower reservoir levels, and damaged crops"
            ]
    # Every emitted anchor and alternative must occur in its pinned source.
    texts: dict[str, str] = {}
    for item in cases:
        for passage in item["expectations"]["passages"]:
            src = next(
                src
                for src in item["sources"]
                if src["alias"] == passage["source_alias"]
            )
            path = FIXTURE_ROOT / src["path"]
            if src["path"] not in texts:
                texts[src["path"]] = normalized(
                    extract_pdf_text(path) if src["kind"] == "pdf" else path.read_text()
                )
            for phrase in [passage["contains"], *passage.get("contains_any", [])]:
                if normalized(phrase) not in texts[src["path"]]:
                    raise ValueError(
                        f"Unverified passage anchor for {item['id']}: {phrase}"
                    )


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output", type=Path, default=ROOT / "evals/cases/real-v1.json"
    )
    args = parser.parse_args()
    download_entries = verify_download_manifest()
    rows, slice_info = write_retail_slice()
    write_rules()
    cases, gold = create_cases(rows)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps({"schema_version": 1, "cases": cases}, ensure_ascii=False, indent=2)
        + "\n",
        encoding="utf-8",
    )
    sources = {}
    for name, entry in download_entries.items():
        path = RAW_ROOT / name
        sources[name] = {
            "path": path.relative_to(FIXTURE_ROOT).as_posix(),
            "sha256": entry["sha256"],
            "bytes": entry["bytes"],
            "url": entry["url"],
            "archive_member": entry.get("archive_member"),
            "attribution": (
                "Chen, D. (2015). Online Retail. UCI Machine Learning Repository."
                if name == "online-retail.xlsx"
                else "Government of India, Economic Survey 2023-24, identified in the source publication."
            ),
            "license": "See the dataset or government publication source page for applicable reuse terms.",
        }
    manifest = {
        "version": VERSION,
        "revision": "2026-10-06",
        "downloads_manifest": {
            "path": "real-v1/downloads.json",
            "sha256": sha256(PACK_ROOT / "downloads.json"),
        },
        "provenance": "Real-source-derived; expected retail values are calculated from the verified original workbook with Decimal. English passage labels are checked against PDF text extraction. Hindi PDF embedded-text limitations and unreviewed language labels are explicitly recorded.",
        "source_files": sources,
        "retail_slice": {
            **slice_info,
            "path": MONTH_CSV.relative_to(FIXTURE_ROOT).as_posix(),
            "sha256": sha256(MONTH_CSV),
            "date_window": "2011-01-01 through 2011-01-31 inclusive",
            "gold": {
                key: (
                    [[country, str(value)] for country, value in value]
                    if key == "top_five_countries"
                    else (
                        {country: str(amount) for country, amount in value.items()}
                        if key == "country_sales"
                        else str(value)
                    )
                )
                for key, value in gold.items()
            },
        },
        "authored_rules": {
            "path": RULES.relative_to(FIXTURE_ROOT).as_posix(),
            "sha256": sha256(RULES),
        },
        "coverage": {
            "total_cases": len(cases),
            "retail_cases": sum("retail" in item["tags"] for item in cases),
            "survey_cases": sum("survey" in item["tags"] for item in cases),
            "document_numeric_calculations": "Narrative numeric claims are not scored as SQL values; review prose and page context manually.",
            "judge_calibration": "not included",
            "multi_turn": "not supported by case schema",
            "OCR": "descriptive/manual only; Hindi embedded-text extraction is garbled in places",
            "language_quality": "requires human review, including Hindi and code-switched cases",
            "automatic_metric_limits": "Top-k category/value lists, chart correctness, clarification quality, passage page number, and prose factuality need manual review.",
        },
    }
    (PACK_ROOT / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    print(
        f"Wrote {len(cases)} cases to {args.output}; retail slice has {len(rows)} rows"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
