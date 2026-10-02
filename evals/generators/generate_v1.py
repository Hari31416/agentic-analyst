#!/usr/bin/env python3
"""Build the deterministic v1 synthetic beneficiary fixtures."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import zipfile
from datetime import datetime
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
import reportlab
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = ROOT / "fixtures" / "v1"
FIXED_ZIP_TIME = (2000, 1, 1, 0, 0, 0)
FIELDS = [
    "application_id",
    "scheme_id",
    "scheme_status",
    "annual_income_inr",
    "grant_amount_inr",
]
ROWS = [
    {
        "application_id": "APP-001",
        "scheme_id": "S1",
        "scheme_status": "active",
        "annual_income_inr": "180000.00",
        "grant_amount_inr": "10000.00",
    },
    {
        "application_id": "APP-002",
        "scheme_id": "S1",
        "scheme_status": "active",
        "annual_income_inr": "200000.00",
        "grant_amount_inr": "15000.00",
    },
    {
        "application_id": "APP-003",
        "scheme_id": "S1",
        "scheme_status": "active",
        "annual_income_inr": "200000.01",
        "grant_amount_inr": "9000.01",
    },
    {
        "application_id": "APP-004",
        "scheme_id": "S1",
        "scheme_status": "inactive",
        "annual_income_inr": "120000.00",
        "grant_amount_inr": "8000.00",
    },
    {
        "application_id": "APP-005",
        "scheme_id": "S2",
        "scheme_status": "active",
        "annual_income_inr": "150000.00",
        "grant_amount_inr": "5000.50",
    },
]
EXPECTED = {
    "version": "v1",
    "currency": "INR",
    "qualifying_rule": "scheme_id = S1 AND scheme_status = active AND annual_income_inr <= 200000.00",
    "threshold_inr": "200000.00",
    "qualifying_application_ids": ["APP-001", "APP-002"],
    "qualifying_count": 2,
    "qualifying_grant_total_inr": "25000.00",
}


def normalize_zip(path: Path) -> None:
    """Rewrite Office ZIP members with fixed timestamps and stable attributes."""
    temp = path.with_suffix(path.suffix + ".normalized")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
        temp, "w", zipfile.ZIP_DEFLATED, compresslevel=9
    ) as dest:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            payload = source.read(name)
            if name == "docProps/core.xml":
                payload = re.sub(
                    rb"(<dcterms:modified\b[^>]*>)[^<]*(</dcterms:modified>)",
                    rb"\g<1>2000-01-01T00:00:00Z\g<2>",
                    payload,
                )
            dest.writestr(info, payload)
    temp.replace(path)


def make_xlsx(output: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "applications"
    sheet.append(FIELDS)
    for row in ROWS:
        sheet.append([row[field] for field in FIELDS])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="245A73")
    for col, width in enumerate([20, 14, 18, 24, 22], start=1):
        sheet.column_dimensions[chr(64 + col)].width = width
    workbook.properties.creator = "Synthetic fixture generator"
    workbook.properties.lastModifiedBy = "Synthetic fixture generator"
    workbook.properties.title = "Synthetic applications v1"
    workbook.properties.subject = "Developer-only synthetic data"
    fixed = datetime(2000, 1, 1, 0, 0, 0)
    workbook.properties.created = fixed
    workbook.properties.modified = fixed
    path = output / "applications.xlsx"
    workbook.save(path)
    normalize_zip(path)


def make_docx(output: Path, language: str) -> None:
    doc = Document()
    if language == "en":
        title, intro = (
            "Synthetic beneficiary fixture (v1)",
            "Developer-only synthetic data; no real people or records are represented.",
        )
        rule = "Qualifies when scheme S1 is active and annual income is at most INR 200,000.00 (inclusive)."
        headings = [
            "Application",
            "Scheme",
            "Status",
            "Annual income (INR)",
            "Grant (INR)",
        ]
        notes = [
            "APP-002 is exactly on the inclusive income threshold.",
            "Amounts are decimal strings with two fractional digits.",
        ]
    else:
        title, intro = (
            "काल्पनिक लाभार्थी नमूना (v1)",
            "केवल डेवलपर के लिए काल्पनिक डेटा; इसमें किसी वास्तविक व्यक्ति या रिकॉर्ड का उपयोग नहीं है।",
        )
        rule = "योजना S1 सक्रिय हो और वार्षिक आय अधिकतम 200000.00 रुपये हो (सीमा शामिल है), तो आवेदन पात्र है।"
        headings = ["आवेदन", "योजना", "स्थिति", "वार्षिक आय (रुपये)", "अनुदान (रुपये)"]
        notes = [
            "APP-002 की आय समावेशी सीमा के बराबर है।",
            "राशियां दो दशमलव अंकों वाले सटीक दशमलव मान हैं।",
        ]
    doc.add_heading(title, 0)
    doc.add_paragraph(intro)
    doc.add_paragraph(rule)
    table = doc.add_table(rows=1, cols=len(headings))
    table.style = "Table Grid"
    for cell, heading in zip(table.rows[0].cells, headings):
        cell.text = heading
    for row in ROWS:
        cells = table.add_row().cells
        for cell, value in zip(cells, [row[field] for field in FIELDS]):
            cell.text = value
    doc.add_heading("Notes" if language == "en" else "टिप्पणियां", level=1)
    for note in notes:
        doc.add_paragraph(note, style="List Bullet")
    doc.core_properties.author = "Synthetic fixture generator"
    doc.core_properties.last_modified_by = "Synthetic fixture generator"
    doc.core_properties.created = datetime(2000, 1, 1)
    doc.core_properties.modified = datetime(2000, 1, 1)
    path = output / f"applications-{language}.docx"
    doc.save(path)
    normalize_zip(path)


def make_pdfs(output: Path) -> None:
    if int(reportlab.Version.split(".", 1)[0]) < 5:
        raise RuntimeError(
            "Hindi PDF generation requires ReportLab 5 or newer for HarfBuzz text shaping"
        )
    try:
        import uharfbuzz  # noqa: F401
    except ImportError as error:
        raise RuntimeError(
            "Hindi PDF generation requires uharfbuzz for HarfBuzz text shaping"
        ) from error
    english = output / "applications-en.pdf"
    c = canvas.Canvas(str(english), pagesize=A4, invariant=1, pageCompression=1)
    c.setTitle("Synthetic applications v1 (English)")
    c.setAuthor("Synthetic fixture generator")
    c.setFont("Helvetica-Bold", 16)
    c.drawString(48, 790, "Synthetic beneficiary fixture (v1)")
    c.setFont("Helvetica", 9)
    c.drawString(
        48,
        772,
        "Developer-only synthetic data; no real people or records are represented.",
    )
    c.drawString(
        48,
        756,
        "Qualifies when scheme S1 is active and annual income is at most INR 200,000.00 (inclusive).",
    )
    headers = ["Application", "Scheme", "Status", "Income INR", "Grant INR"]
    positions = [48, 153, 235, 315, 411]
    y = 720
    c.setFont("Helvetica-Bold", 8)
    for x, value in zip(positions, headers):
        c.drawString(x, y, value)
    c.setFont("Helvetica", 8)
    for row in ROWS:
        y -= 20
        for x, field in zip(positions, FIELDS):
            c.drawString(x, y, row[field])
    c.drawString(
        48,
        y - 35,
        "APP-002 is exactly on the inclusive threshold. Decimal values retain two places.",
    )
    c.save()

    hindi = output / "applications-hi.pdf"
    pdfmetrics.registerFont(
        TTFont(
            "NotoDevanagari",
            str(ROOT / "fonts" / "NotoSansDevanagari-Regular.ttf"),
            shapable=True,
        )
    )
    pdfmetrics.registerFont(
        TTFont("NotoLatin", str(ROOT / "fonts" / "NotoSans-Regular.ttf"), shapable=True)
    )
    c = canvas.Canvas(str(hindi), pagesize=A4, invariant=1, pageCompression=1)
    c.setTitle("Synthetic applications v1 (Hindi)")
    c.setAuthor("Synthetic fixture generator")
    c.setFillColorRGB(0.09, 0.17, 0.21)
    c.setFont("NotoDevanagari", 18)
    c.drawString(48, 785, "काल्पनिक लाभार्थी नमूना", shaping=True)
    c.setFont("NotoLatin", 10)
    c.drawString(48, 768, "v1", shaping=True)
    c.setFont("NotoDevanagari", 10)
    c.drawString(
        48,
        760,
        "केवल डेवलपर के लिए काल्पनिक डेटा; इसमें किसी वास्तविक व्यक्ति या रिकॉर्ड का उपयोग नहीं है।",
        shaping=True,
    )
    c.drawString(
        48,
        738,
        "योजना संख्या एक सक्रिय हो और वार्षिक आय अधिकतम 200000.00 रुपये हो (सीमा शामिल है), तो आवेदन पात्र है।",
        shaping=True,
    )
    c.drawString(
        48,
        710,
        "आवेदन | योजना | स्थिति | वार्षिक आय (रुपये) | अनुदान (रुपये)",
        shaping=True,
    )
    c.setFont("NotoLatin", 8)
    x_positions = [48, 150, 225, 294, 390]
    for x, text in zip(
        x_positions, ["Application", "Scheme", "Status", "Income INR", "Grant INR"]
    ):
        c.drawString(x, 690, text, shaping=True)
    for index, row in enumerate(ROWS):
        y = 670 - index * 18
        for x, field in zip(x_positions, FIELDS):
            c.drawString(x, y, row[field], shaping=True)
    c.setFont("NotoDevanagari", 10)
    c.drawString(48, 565, "दूसरे आवेदन की आय समावेशी सीमा के बराबर है।", shaping=True)
    c.drawString(
        48, 545, "राशियां दो दशमलव अंकों वाले सटीक दशमलव मान हैं।", shaping=True
    )
    c.save()


def make_sql(output: Path) -> None:
    columns = "application_id VARCHAR(20) PRIMARY KEY, scheme_id VARCHAR(10) NOT NULL, scheme_status VARCHAR(10) NOT NULL, annual_income_inr DECIMAL(12,2) NOT NULL, grant_amount_inr DECIMAL(12,2) NOT NULL"
    inserts = ",\n".join(
        "('"
        + "', '".join(row[field] for field in FIELDS[:3])
        + "', "
        + row["annual_income_inr"]
        + ", "
        + row["grant_amount_inr"]
        + ")"
        for row in ROWS
    )
    for dialect, create in [
        (
            "postgresql",
            f"CREATE TABLE IF NOT EXISTS synthetic_applications ({columns});",
        ),
        (
            "mysql",
            f"CREATE TABLE IF NOT EXISTS synthetic_applications ({columns}) ENGINE=InnoDB;",
        ),
    ]:
        text = (
            "-- DEVELOPER-ONLY synthetic benchmark seed; never use in production.\n"
            "-- Idempotently replaces only this fixture-owned table's rows.\n"
            f"{create}\nDELETE FROM synthetic_applications WHERE application_id LIKE 'APP-%';\n"
            "INSERT INTO synthetic_applications (application_id, scheme_id, scheme_status, annual_income_inr, grant_amount_inr) VALUES\n"
            f"{inserts};\n"
        )
        (output / f"seed-{dialect}.sql").write_text(
            text, encoding="utf-8", newline="\n"
        )


def make_data(output: Path) -> None:
    with (output / "applications.csv").open("w", encoding="utf-8", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=FIELDS, lineterminator="\n")
        writer.writeheader()
        writer.writerows(ROWS)
    make_xlsx(output)
    make_sql(output)
    make_docx(output, "en")
    make_docx(output, "hi")
    make_pdfs(output)
    (output / "expected.json").write_text(
        json.dumps(EXPECTED, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
        newline="\n",
    )
    hashes = {
        p.name: hashlib.sha256(p.read_bytes()).hexdigest()
        for p in sorted(output.iterdir())
        if p.is_file() and p.name != "manifest.json"
    }
    font_inputs = {
        p.name: hashlib.sha256(p.read_bytes()).hexdigest()
        for p in sorted((ROOT / "fonts").glob("*"))
        if p.is_file()
    }
    manifest = {
        "version": "v1",
        "algorithm": "sha256",
        "files": hashes,
        "font_inputs": font_inputs,
    }
    (output / "manifest.json").write_text(
        json.dumps(manifest, indent=2) + "\n", encoding="utf-8", newline="\n"
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    output = args.output.resolve()
    output.mkdir(parents=True, exist_ok=True)
    make_data(output)
    print(f"Generated {len(list(output.iterdir()))} files in {output}")


if __name__ == "__main__":
    main()
