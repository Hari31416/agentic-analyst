import csv
import hashlib
import json
import subprocess
import sys
from decimal import Decimal
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

EVALS = Path(__file__).resolve().parents[1]
FIXTURES = EVALS / "fixtures" / "v1"


def test_expected_result_matches_csv_decimal_arithmetic():
    with (FIXTURES / "applications.csv").open(encoding="utf-8", newline="") as stream:
        rows = list(csv.DictReader(stream))
    eligible = [
        row
        for row in rows
        if row["scheme_id"] == "S1"
        and row["scheme_status"] == "active"
        and Decimal(row["annual_income_inr"]) <= Decimal("200000.00")
    ]
    expected = json.loads((FIXTURES / "expected.json").read_text(encoding="utf-8"))
    assert [row["application_id"] for row in eligible] == expected[
        "qualifying_application_ids"
    ]
    assert len(eligible) == expected["qualifying_count"] == 2
    assert sum(
        (Decimal(row["grant_amount_inr"]) for row in eligible), Decimal("0.00")
    ) == Decimal(expected["qualifying_grant_total_inr"])


def test_csv_xlsx_and_seed_sql_have_the_same_five_rows():
    with (FIXTURES / "applications.csv").open(encoding="utf-8", newline="") as stream:
        csv_rows = list(csv.DictReader(stream))
    sheet = load_workbook(FIXTURES / "applications.xlsx", data_only=True).active
    xlsx_rows = [
        dict(zip([cell.value for cell in sheet[1]], values))
        for values in sheet.iter_rows(min_row=2, values_only=True)
    ]
    assert xlsx_rows == csv_rows
    for dialect in ("mysql", "postgresql"):
        sql = (FIXTURES / f"seed-{dialect}.sql").read_text(encoding="utf-8")
        assert "DEVELOPER-ONLY" in sql
        assert sql.count("('APP-") == 5
        for row in csv_rows:
            assert (
                f"('{row['application_id']}', '{row['scheme_id']}', '{row['scheme_status']}', {row['annual_income_inr']}, {row['grant_amount_inr']})"
                in sql
            )


def test_english_and_hindi_documents_contain_original_language_text():
    en = "\n".join(
        p.text for p in Document(FIXTURES / "applications-en.docx").paragraphs
    )
    hi = "\n".join(
        p.text for p in Document(FIXTURES / "applications-hi.docx").paragraphs
    )
    assert "at most INR 200,000.00" in en
    assert "योजना S1 सक्रिय" in hi
    for language in ("en", "hi"):
        pdf = PdfReader(FIXTURES / f"applications-{language}.pdf")
        assert len(pdf.pages) == 1
        assert "Synthetic applications v1" in pdf.metadata.title
    hindi_pdf_text = PdfReader(FIXTURES / "applications-hi.pdf").pages[0].extract_text()
    assert "काल्पनिक लाभार्थी नमूना" in hindi_pdf_text
    assert "APP-002" in hindi_pdf_text


def test_manifest_covers_every_fixture_payload():
    manifest = json.loads((FIXTURES / "manifest.json").read_text(encoding="utf-8"))
    actual = {
        path.name: hashlib.sha256(path.read_bytes()).hexdigest()
        for path in sorted(FIXTURES.iterdir())
        if path.is_file() and path.name != "manifest.json"
    }
    assert manifest["algorithm"] == "sha256"
    assert manifest["files"] == actual
    expected_fonts = {
        path.name: hashlib.sha256(path.read_bytes()).hexdigest()
        for path in sorted((EVALS / "fonts").iterdir())
        if path.is_file()
    }
    assert manifest["font_inputs"] == expected_fonts


def test_regeneration_is_byte_reproducible(tmp_path):
    output = tmp_path / "v1"
    subprocess.run(
        [
            sys.executable,
            str(EVALS / "generators" / "generate_v1.py"),
            "--output",
            str(output),
        ],
        check=True,
        capture_output=True,
        text=True,
    )
    for source in FIXTURES.iterdir():
        if source.is_file():
            assert (
                output / source.name
            ).read_bytes() == source.read_bytes(), source.name
